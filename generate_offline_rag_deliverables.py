import html
import os
import sys
import zipfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).parent / ".vendor"))

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image, ImageDraw, ImageFont


OUT = Path("交付物")
IMG = OUT / "images"
OUT.mkdir(exist_ok=True)
IMG.mkdir(exist_ok=True)

FONT_MED = "/System/Library/Fonts/STHeiti Medium.ttc"
FONT_LIGHT = "/System/Library/Fonts/STHeiti Light.ttc"


def font(size, bold=False):
    return ImageFont.truetype(FONT_MED if bold else FONT_LIGHT, size)


def box(d, xy, title, lines, fill, outline):
    x1, y1, x2, y2 = xy
    d.rounded_rectangle(xy, radius=18, fill=fill, outline=outline, width=3)
    d.text((x1 + 18, y1 + 14), title, font=font(24, True), fill="#111827")
    y = y1 + 56
    for line in lines:
        d.text((x1 + 18, y), line, font=font(18), fill="#334155")
        y += 27


def arrow(d, start, end, color="#64748B", width=4):
    d.line([start, end], fill=color, width=width)
    import math
    x1, y1 = start
    x2, y2 = end
    ang = math.atan2(y2 - y1, x2 - x1)
    l = 15
    p1 = (x2 + l * math.cos(ang + 2.45), y2 + l * math.sin(ang + 2.45))
    p2 = (x2 + l * math.cos(ang - 2.45), y2 + l * math.sin(ang - 2.45))
    d.polygon([(x2, y2), p1, p2], fill=color)


def architecture_png():
    img = Image.new("RGB", (1800, 1120), "#FFFFFF")
    d = ImageDraw.Draw(img)
    d.text((60, 42), "离线 RAG Agent 总体技术架构", font=font(44, True), fill="#0F172A")
    d.text((60, 100), "应用层 → 接入层 → Agent编排层 → Tool/RAG层 → 模型适配层 → 数据存储层", font=font(24), fill="#475569")
    layers = [
        ("应用层", ["Web管理台", "Java业务系统", "字段填充API", "知识库管理"], "#DBEAFE", "#2563EB"),
        ("接入层", ["Nginx/APISIX", "鉴权", "限流", "审计入口"], "#E0F2FE", "#0284C7"),
        ("Agent编排层", ["FastAPI", "LangGraph 1.x", "任务状态机", "人工确认"], "#DCFCE7", "#16A34A"),
        ("Tool / RAG层", ["MCP/Tool", "字段字典", "精确查询", "混合检索/Rerank"], "#FEF3C7", "#D97706"),
        ("模型适配层", ["CPU llama.cpp", "API候选模型", "GPU vLLM", "Embedding/Rerank"], "#F3E8FF", "#7C3AED"),
        ("数据存储层", ["PostgreSQL 16", "pgvector", "excel_rows", "rag_chunks/audit"], "#FFE4E6", "#E11D48"),
    ]
    y = 170
    for name, items, fill, outline in layers:
        d.rounded_rectangle((70, y, 1730, y + 118), radius=22, fill=fill, outline=outline, width=3)
        d.text((105, y + 38), name, font=font(30, True), fill="#111827")
        x = 370
        for item in items:
            d.rounded_rectangle((x, y + 30, x + 285, y + 88), radius=16, fill="#FFFFFF", outline=outline, width=2)
            d.text((x + 18, y + 47), item, font=font(21), fill="#1F2937")
            x += 320
        if y < 170 + 5 * 145:
            arrow(d, (900, y + 118), (900, y + 145), outline, 5)
        y += 145
    d.rounded_rectangle((70, 1040, 1730, 1090), radius=16, fill="#F8FAFC", outline="#CBD5E1", width=2)
    d.text((100, 1053), "横切能力：OpenTelemetry + Prometheus + Grafana + Phoenix；Presidio 脱敏 + PostgreSQL 审计", font=font(22), fill="#334155")
    p = IMG / "overall_architecture.png"
    img.save(p)
    return p


def pipeline_png():
    img = Image.new("RGB", (1800, 1100), "#FFFFFF")
    d = ImageDraw.Draw(img)
    d.text((60, 42), "离线构建与在线查询双 Pipeline", font=font(44, True), fill="#0F172A")
    d.text((60, 100), "Excel 以“一行=主chunk、row_json保存整行”为核心；查询时结构化精确查询优先，语义召回补充。", font=font(24), fill="#475569")
    d.rounded_rectangle((50, 150, 1750, 555), radius=28, fill="#F8FAFC", outline="#16A34A", width=4)
    d.text((80, 178), "A. 离线构建 / 写入 RAG", font=font(30, True), fill="#166534")
    stages = [
        ("文件输入", ["Excel/CSV", "Word/PPT/PDF", "TXT/Markdown"]),
        ("解析标准化", ["结构化解析", "页码/标题/OCR", "字段归一化"]),
        ("分片策略", ["Excel: 一行=chunk", "row_json整行保存", "长文本二级分片"]),
        ("向量与索引", ["bge-m3 embedding", "B-tree/GIN/trgm", "pgvector HNSW"]),
        ("质量回滚", ["批次状态", "抽样召回", "失败回滚"]),
    ]
    x = 90
    for i, (t, lines) in enumerate(stages):
        box(d, (x, 250, x + 270, 465), t, lines, "#ECFDF5", "#16A34A")
        if i < 4:
            arrow(d, (x + 270, 357), (x + 325, 357), "#16A34A")
        x += 335
    d.rounded_rectangle((50, 625, 1750, 1030), radius=28, fill="#F8FAFC", outline="#7C3AED", width=4)
    d.text((80, 653), "B. 在线查询 / 字段填充 RAG", font=font(30, True), fill="#5B21B6")
    stages = [
        ("用户问题", ["字段填充", "事实问答", "模糊查询"]),
        ("查询理解", ["意图识别", "字段名/值抽取", "字段别名映射"]),
        ("检索路由", ["字段值→精确查", "自然语言→向量", "混合→过滤召回"]),
        ("结果融合", ["row_id优先", "TopK重排", "置信度评分"]),
        ("答案输出", ["字段值直返", "LLM解释", "引用/审计"]),
    ]
    x = 90
    for i, (t, lines) in enumerate(stages):
        box(d, (x, 725, x + 270, 940), t, lines, "#F5F3FF", "#7C3AED")
        if i < 4:
            arrow(d, (x + 270, 832), (x + 325, 832), "#7C3AED")
        x += 335
    p = IMG / "pipeline.png"
    img.save(p)
    return p


def sequence_png():
    img = Image.new("RGB", (1800, 1000), "#FFFFFF")
    d = ImageDraw.Draw(img)
    d.text((60, 42), "字段填充查询时序图", font=font(44, True), fill="#0F172A")
    actors = ["用户/Java系统", "API网关", "FastAPI", "LangGraph", "RAG服务", "PostgreSQL", "LLM Router"]
    xs = [120, 360, 600, 860, 1120, 1380, 1620]
    for x, a in zip(xs, actors):
        d.rounded_rectangle((x - 95, 150, x + 95, 210), radius=14, fill="#E0F2FE", outline="#0284C7", width=2)
        d.text((x - 74, 168), a, font=font(18, True), fill="#0F172A")
        d.line((x, 220, x, 900), fill="#CBD5E1", width=3)
    steps = [
        (0, 1, "1 提交字段填充/问答请求"),
        (1, 2, "2 鉴权、限流、数据域过滤"),
        (2, 3, "3 脱敏、创建任务状态"),
        (3, 4, "4 Tool: 查询理解与检索路由"),
        (4, 5, "5a 明确字段值：精确查询 excel_rows"),
        (4, 5, "5b 自然语言：pgvector 语义召回"),
        (4, 3, "6 返回 row_id/chunk/source"),
        (3, 6, "7 复杂问题调用 LLM"),
        (6, 3, "8 JSON Schema 约束输出"),
        (3, 2, "9 融合、置信度、审计"),
        (2, 0, "10 返回字段值/答案/引用"),
    ]
    y = 250
    for a, b, label in steps:
        color = "#16A34A" if "5a" in label else "#7C3AED" if "5b" in label or "LLM" in label else "#475569"
        arrow(d, (xs[a], y), (xs[b], y), color, 4)
        d.text((min(xs[a], xs[b]) + 12, y - 28), label, font=font(17), fill="#111827")
        y += 58
    p = IMG / "sequence.png"
    img.save(p)
    return p


def deployment_png():
    img = Image.new("RGB", (1800, 980), "#FFFFFF")
    d = ImageDraw.Draw(img)
    d.text((60, 42), "单机虚拟机离线部署架构", font=font(44, True), fill="#0F172A")
    d.text((60, 100), "前期 CPU 验证功能，确认模型效果后再选择 GPU；所有依赖以离线包交付。", font=font(24), fill="#475569")
    d.rounded_rectangle((70, 160, 1730, 870), radius=28, fill="#F8FAFC", outline="#334155", width=4)
    comps = [
        ((120, 270, 430, 410), "Nginx/APISIX", ["80/443", "鉴权/限流"], "#DBEAFE", "#2563EB"),
        ((520, 270, 850, 410), "rag-agent-api", ["FastAPI", "LangGraph 1.x"], "#DCFCE7", "#16A34A"),
        ((940, 270, 1270, 410), "worker", ["解析/embedding", "批处理/重试"], "#FEF3C7", "#D97706"),
        ((1360, 270, 1660, 410), "llm-router", ["CPU/API/GPU", "统一接口"], "#F3E8FF", "#7C3AED"),
        ((120, 520, 430, 680), "PostgreSQL 16", ["pgvector/JSONB", "audit/rag/excel"], "#FFE4E6", "#E11D48"),
        ((520, 520, 850, 680), "Redis 7", ["缓存/队列", "限流计数"], "#FFEDD5", "#EA580C"),
        ((940, 520, 1270, 680), "文件与模型目录", ["原始文件", "模型/wheel/镜像"], "#E0F2FE", "#0284C7"),
        ((1360, 520, 1660, 680), "可观测", ["Prometheus", "Grafana/Phoenix"], "#ECFDF5", "#059669"),
    ]
    for xy, t, lines, fill, outline in comps:
        box(d, xy, t, lines, fill, outline)
    for s, e in [((430, 340), (520, 340)), ((850, 340), (940, 340)), ((1270, 340), (1360, 340)), ((685, 410), (685, 520)), ((1105, 410), (1105, 520))]:
        arrow(d, s, e, "#475569", 4)
    d.text((120, 770), "CPU验证最低：16C / 64GB / 500GB SSD；功能跑通，不以模型延迟作为性能指标。", font=font(24), fill="#334155")
    d.text((120, 812), "GPU生产建议：32C / 128GB / 2TB SSD + L20 48GB 或 A100 40GB；vLLM承载14B/32B量化模型。", font=font(24), fill="#334155")
    p = IMG / "deployment.png"
    img.save(p)
    return p


def shade(cell, color):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), color)
    tc_pr.append(shd)


def setup_doc(doc):
    sec = doc.sections[0]
    sec.top_margin = Inches(0.65)
    sec.bottom_margin = Inches(0.65)
    sec.left_margin = Inches(0.72)
    sec.right_margin = Inches(0.72)
    for style in ["Normal", "Heading 1", "Heading 2", "Heading 3"]:
        doc.styles[style].font.name = "Arial Unicode MS"
    doc.styles["Normal"].font.size = Pt(10.5)
    doc.styles["Heading 1"].font.size = Pt(18)
    doc.styles["Heading 2"].font.size = Pt(14)


def title(doc, main, sub):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(main)
    r.bold = True
    r.font.size = Pt(22)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(sub)
    r.font.size = Pt(11)
    r.font.color.rgb = RGBColor(71, 85, 105)


def bullets(doc, items):
    for item in items:
        doc.add_paragraph(item, style="List Bullet")


def table(doc, headers, rows):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(headers):
        c = t.rows[0].cells[i]
        c.text = h
        c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        shade(c, "E2E8F0")
        for p in c.paragraphs:
            for r in p.runs:
                r.bold = True
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = str(v)
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
    doc.add_paragraph("")


def picture(doc, path, cap):
    doc.add_picture(str(path), width=Inches(6.85))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = doc.add_paragraph(cap)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for r in p.runs:
        r.font.size = Pt(9)
        r.font.color.rgb = RGBColor(100, 116, 139)


def build_tech(paths):
    doc = Document()
    setup_doc(doc)
    title(doc, "离线 RAG Agent 技术选型文档", "Excel 行级字段填充 + 多格式知识库 + PostgreSQL/pgvector + 内网离线部署")
    doc.add_heading("1. 需求复述", 1)
    bullets(doc, [
        "需要基于现有离线 RAG 场景设计全新的 Agent 架构，覆盖离线构建和在线查询两条链路。",
        "文件来源包括 Excel/CSV、Word、PPT、PDF、TXT/Markdown/HTML，其中 Excel 是核心：一行代表一条业务记录，用户可能根据某个字段值填充同一行其他字段。",
        "知识库规模不大，字段量级几万，优先沿用 PostgreSQL + pgvector；当前阶段不建议优先引入 Elasticsearch。",
        "公司后端主栈为 Java，但 AI/RAG/Agent 核心建议使用 Python；Java 保留业务接入、权限、管理后台与企业集成。",
        "部署在内网单机虚机，不能访问外网，需要离线包；前期 CPU 只验证功能，模型确认后再选择 GPU。",
    ])
    doc.add_heading("2. 选型总览", 1)
    table(doc, ["类别", "组件与版本", "说明"], [
        ("语言", "Python 3.11.9 + Java 17", "Python 承载 AI 核心，Java 承载企业业务接入。"),
        ("Agent", "LangGraph >= 1.0.3，LangChain >= 1.0.5", "1.x API 稳定，适合状态机编排、断点恢复和可观测。"),
        ("API", "FastAPI 0.115.x + Uvicorn 0.32.x", "轻量、异步、Pydantic 2.x 校验，支持 SSE。"),
        ("任务", "Celery 5.4.x / RQ 2.x + Redis 7.4.x", "离线解析、向量化、重试、缓存。"),
        ("存储", "PostgreSQL 16.x + pgvector 0.8.x", "普通列/JSONB/pg_trgm/向量一体，适合当前规模。"),
        ("文档解析", "openpyxl 3.1.x、pandas 2.2.x、python-docx 1.1.x、PyMuPDF 1.24.x、python-pptx 1.0.x", "覆盖 Excel、Word、PDF、PPT 与文本。"),
        ("Embedding/Rerank", "bge-m3 + bge-reranker-v2-m3", "中文效果稳定，适合混合检索和重排。"),
        ("模型路由", "CPU llama.cpp；API 百炼/方舟；GPU vLLM 0.8.x+", "先功能验证，再做模型和 GPU 选型。"),
        ("审计脱敏", "Presidio 2.2.x + 自定义正则 + audit schema", "覆盖手机号、身份证、邮箱、合同号、客户名等。"),
        ("可观测", "OpenTelemetry 1.30.x + Prometheus + Grafana + Phoenix", "LLM/Tool/DB 调用可追踪。"),
    ])
    doc.add_heading("3. Python 与 Java 对比", 1)
    table(doc, ["维度", "Java", "Python", "结论"], [
        ("AI 生态", "Spring AI、LangChain4j 可用但能力滞后", "LangChain/LangGraph/vLLM/Transformers 首发支持", "Agent 核心选 Python"),
        ("企业集成", "权限、网关、后台、治理成熟", "需要补工程规范", "业务接入选 Java"),
        ("文档解析", "POI/Tika 稳定", "pandas/openpyxl/PyMuPDF/unstructured 灵活", "解析层选 Python 更快"),
        ("模型推理", "通常调用外部服务", "vLLM/llama.cpp/Embedding 生态完整", "模型服务选 Python"),
        ("综合", "适合外围系统", "适合 AI 核心", "Java + Python 分层解耦"),
    ])
    doc.add_heading("4. pgvector 评估", 1)
    bullets(doc, [
        "当前主要是 Excel 字段查询和行级补全，PostgreSQL 的 B-tree、JSONB GIN、pg_trgm 和 pgvector 组合足够。",
        "Excel 行数据必须单独落 excel_rows，不能只依赖向量库；向量召回用于自然语言语义补充。",
        "只有当大量非结构化全文检索、复杂中文分词、高亮、聚合统计、高并发搜索成为刚需时，再引入 Elasticsearch/OpenSearch。",
    ])
    picture(doc, paths["pipeline"], "图 1：离线构建与在线查询双 Pipeline")
    doc.add_heading("5. 最低环境配置", 1)
    table(doc, ["阶段", "最低配置", "用途", "说明"], [
        ("CPU 功能验证", "16C / 64GB / 500GB SSD", "跑通上传、解析、入库、检索、Agent、审计", "7B GGUF 可慢速验证，不以延迟作为性能指标。"),
        ("CPU 增强验证", "32C / 128GB / 1TB SSD", "批量导入与并发联调", "适合验证非模型环节 100 QPS。"),
        ("GPU 生产起步", "32C / 128GB / 2TB SSD + L20 48GB 或 A100 40GB", "vLLM 推理 14B/32B 量化模型", "确认模型效果后再采购。"),
        ("离线交付", "Docker 镜像 tar + wheelhouse + 模型权重 + SQL", "内网无外网部署", "版本锁定并附 SHA256。"),
    ])
    doc.add_heading("6. 参考", 1)
    bullets(doc, ["LangChain 官方文档：https://docs.langchain.com/", "LangGraph 官方文档：https://docs.langchain.com/oss/python/langgraph/overview", "pgvector：https://github.com/pgvector/pgvector"])
    doc.save(OUT / "离线RAG智能体_技术选型文档.docx")


def build_arch(paths):
    doc = Document()
    setup_doc(doc)
    title(doc, "离线 RAG Agent 架构设计文档", "总体架构、处理流程、时序图、部署图、Tool/MCP/Skill 交互细节")
    doc.add_heading("1. 需求理解", 1)
    bullets(doc, [
        "系统需要处理两类任务：离线知识构建和在线查询/字段填充。",
        "离线构建把多格式文件解析成标准 DocumentModel，再按规则分片、向量化、写入 PostgreSQL/pgvector。",
        "在线查询先做鉴权和查询理解，再根据“字段明确程度”选择精确查询、向量检索或混合检索。",
        "字段填充场景应尽量不调用 LLM，直接基于 row_id 返回同一行字段，减少成本和幻觉。",
    ])
    doc.add_heading("2. 总体架构", 1)
    picture(doc, paths["arch"], "图 1：离线 RAG Agent 总体技术架构")
    doc.add_heading("3. 处理流程", 1)
    picture(doc, paths["pipeline"], "图 2：离线构建与在线查询处理流程")
    doc.add_heading("4. 时序图", 1)
    picture(doc, paths["sequence"], "图 3：字段填充查询时序图")
    doc.add_heading("5. 模块设计", 1)
    table(doc, ["模块", "职责", "关键逻辑"], [
        ("文件接入模块", "文件上传、目录扫描、批次管理", "生成 source_id、batch_id、checksum，记录原始文件路径。"),
        ("解析模块", "按文件类型解析", "Excel 结构化解析；Word/PPT/PDF 按标题、页码、slide、OCR 解析。"),
        ("分片模块", "生成 chunk 与 metadata", "Excel 一行一个主 chunk，row_json 保留整行；非 Excel 按标题/段落/页切分。"),
        ("索引模块", "写入 PG 与向量库", "rag_chunks 保存文本和向量，excel_rows 保存整行数据，field_dictionary 保存字段别名。"),
        ("查询理解模块", "识别意图与字段", "抽取字段名、字段值、目标字段，完成别名映射。"),
        ("检索路由模块", "选择查询路径", "明确字段值走精确查询；自然语言走向量；混合场景先过滤再召回。"),
        ("Agent 编排模块", "串联状态与工具", "LangGraph 管理状态机、重试、人工确认、审计。"),
        ("LLM Router", "统一模型入口", "CPU/API/GPU 三模式切换，对上层暴露统一接口。"),
    ])
    doc.add_heading("6. Tool / MCP / Skill 交互", 1)
    bullets(doc, [
        "Tool：字段字典查询、PostgreSQL 精确查询、pgvector 检索、Rerank、审计写入、脱敏均封装成确定性工具。",
        "MCP：后续可把数据库、文件系统、内部知识源封装为 MCP Server，供不同 Agent 复用。",
        "Skill：把 Excel 行级分片、字段别名学习、质量检查等专家流程沉淀为稳定技能，避免每次依赖 LLM 临场发挥。",
        "Agent：只负责决策和编排，不负责执行可确定的脏活；确定性逻辑全部下沉到工具和技能。",
    ])
    doc.add_heading("7. 部署架构", 1)
    picture(doc, paths["deployment"], "图 4：单机虚拟机离线部署架构")
    doc.add_heading("8. 高性能策略", 1)
    bullets(doc, [
        "大 JSON 不直接进入模型：先入库，再按字段、表、业务域分批处理。",
        "结构化优先：字段填充优先 PostgreSQL 精确查询，必要时再语义召回。",
        "缓存：schema hash、字段别名、query embedding、常见查询结果进入 Redis。",
        "异步化：解析、embedding、批量导入走 worker 队列，可重试可回滚。",
        "限流：网关限流 + LLM Router 并发控制，防止模型服务被打满。",
    ])
    doc.save(OUT / "离线RAG智能体_架构设计文档.docx")


def pptx_escape(s):
    return html.escape(str(s), quote=True)


def sp_text(idx, x, y, w, h, text, size=2400, bold=False, fill="FFFFFF", line="CBD5E1", color="111827"):
    b = " b=\"1\"" if bold else ""
    return f"""
    <p:sp><p:nvSpPr><p:cNvPr id=\"{idx}\" name=\"Text {idx}\"/><p:cNvSpPr/><p:nvPr/></p:nvSpPr>
      <p:spPr><a:xfrm><a:off x=\"{x}\" y=\"{y}\"/><a:ext cx=\"{w}\" cy=\"{h}\"/></a:xfrm><a:prstGeom prst=\"roundRect\"><a:avLst/></a:prstGeom><a:solidFill><a:srgbClr val=\"{fill}\"/></a:solidFill><a:ln w=\"16000\"><a:solidFill><a:srgbClr val=\"{line}\"/></a:solidFill></a:ln></p:spPr>
      <p:txBody><a:bodyPr wrap=\"square\"/><a:lstStyle/><a:p><a:r><a:rPr lang=\"zh-CN\" sz=\"{size}\"{b}><a:solidFill><a:srgbClr val=\"{color}\"/></a:solidFill></a:rPr><a:t>{pptx_escape(text)}</a:t></a:r></a:p></p:txBody>
    </p:sp>"""


def pic(idx, rid, x, y, w, h):
    return f"""
    <p:pic><p:nvPicPr><p:cNvPr id=\"{idx}\" name=\"Picture {idx}\"/><p:cNvPicPr/><p:nvPr/></p:nvPicPr>
      <p:blipFill><a:blip r:embed=\"{rid}\"/><a:stretch><a:fillRect/></a:stretch></p:blipFill>
      <p:spPr><a:xfrm><a:off x=\"{x}\" y=\"{y}\"/><a:ext cx=\"{w}\" cy=\"{h}\"/></a:xfrm><a:prstGeom prst=\"rect\"><a:avLst/></a:prstGeom></p:spPr>
    </p:pic>"""


def slide_xml(shapes):
    return f"""<?xml version=\"1.0\" encoding=\"UTF-8\" standalone=\"yes\"?>
<p:sld xmlns:a=\"http://schemas.openxmlformats.org/drawingml/2006/main\" xmlns:r=\"http://schemas.openxmlformats.org/officeDocument/2006/relationships\" xmlns:p=\"http://schemas.openxmlformats.org/presentationml/2006/main\">
  <p:cSld><p:bg><p:bgPr><a:solidFill><a:srgbClr val=\"FFFFFF\"/></a:solidFill></p:bgPr></p:bg><p:spTree>
    <p:nvGrpSpPr><p:cNvPr id=\"1\" name=\"\"/><p:cNvGrpSpPr/><p:nvPr/></p:nvGrpSpPr><p:grpSpPr><a:xfrm><a:off x=\"0\" y=\"0\"/><a:ext cx=\"0\" cy=\"0\"/><a:chOff x=\"0\" y=\"0\"/><a:chExt cx=\"0\" cy=\"0\"/></a:xfrm></p:grpSpPr>
    {''.join(shapes)}
  </p:spTree></p:cSld><p:clrMapOvr><a:masterClrMapping/></p:clrMapOvr>
</p:sld>"""


def make_pptx(paths):
    ppt = OUT / "离线RAG智能体_架构与选型汇报.pptx"
    slides = []
    rels = []
    title_shapes = [
        sp_text(2, 700000, 1200000, 10500000, 900000, "离线 RAG Agent 架构与技术选型", 3800, True, "0F172A", "0F172A", "FFFFFF"),
        sp_text(3, 760000, 2200000, 10800000, 520000, "Excel 行级字段填充 · PostgreSQL + pgvector · LangGraph 1.x · 内网离线部署", 1800, False, "0F172A", "0F172A", "CBD5E1"),
    ]
    slides.append(slide_xml(title_shapes)); rels.append("")
    slides.append(slide_xml([
        sp_text(2, 500000, 350000, 11200000, 520000, "需求复述", 3000, True, "FFFFFF", "FFFFFF"),
        sp_text(3, 800000, 1300000, 11000000, 4200000, "1. 离线构建：多格式文件解析、Excel 行级 chunk、向量化与写入。\n2. 在线查询：字段值精确过滤、语义召回、结果融合、字段填充。\n3. 内网单机虚机部署，不能访问外网，必须离线交付。\n4. 前期 CPU 验证功能，确认模型后再选择 GPU。\n5. 需要审计、脱敏、可观测、人工确认与错误样本回流。", 1800, False, "F8FAFC", "CBD5E1"),
    ])); rels.append("")
    image_slides = [
        ("总体技术架构", paths["arch"]),
        ("离线构建与在线查询流程", paths["pipeline"]),
        ("字段填充查询时序", paths["sequence"]),
        ("单机虚拟机部署架构", paths["deployment"]),
    ]
    for i, (title_txt, image_path) in enumerate(image_slides, start=3):
        slides.append(slide_xml([
            sp_text(2, 500000, 260000, 11500000, 450000, title_txt, 2600, True, "FFFFFF", "FFFFFF"),
            pic(3, "rId1", 500000, 900000, 11280000, 5400000),
        ]))
        rels.append(f"""<?xml version=\"1.0\" encoding=\"UTF-8\" standalone=\"yes\"?>
<Relationships xmlns=\"http://schemas.openxmlformats.org/package/2006/relationships\"><Relationship Id=\"rId1\" Type=\"http://schemas.openxmlformats.org/officeDocument/2006/relationships/image\" Target=\"../media/{image_path.name}\"/></Relationships>""")
    slides.append(slide_xml([
        sp_text(2, 500000, 350000, 11200000, 520000, "技术选型结论", 3000, True, "FFFFFF", "FFFFFF"),
        sp_text(3, 800000, 1300000, 11000000, 4200000, "• AI 核心：Python 3.11 + FastAPI + LangGraph 1.x + LangChain 1.x。\n• 业务接入：Java 17 / Spring Boot 3.3.x。\n• 存储：PostgreSQL 16 + pgvector 0.8，结构化字段优先。\n• 模型：CPU llama.cpp + 7B GGUF；GPU vLLM + 14B/32B 量化模型。\n• 离线交付：Docker 镜像、wheelhouse、模型权重、初始化 SQL。", 1800, False, "F8FAFC", "CBD5E1"),
    ])); rels.append("")
    slides.append(slide_xml([
        sp_text(2, 500000, 350000, 11200000, 520000, "为什么第一阶段不引入 Elasticsearch", 3000, True, "FFFFFF", "FFFFFF"),
        sp_text(3, 800000, 1300000, 11000000, 4200000, "• 当前核心是 Excel 指定字段查找与同一行字段补全，不是大规模全文搜索。\n• PostgreSQL 已覆盖 B-tree、JSONB GIN、pg_trgm、全文检索和 pgvector。\n• 少一个组件，离线部署和运维复杂度显著下降。\n• 后续出现复杂中文分词、高亮、聚合统计、高并发全文搜索时，再引入 ES/OpenSearch。", 1800, False, "F8FAFC", "CBD5E1"),
    ])); rels.append("")

    template = Path(".vendor/pptx/templates/default.pptx")
    with zipfile.ZipFile(template) as zin:
        base_files = {n: zin.read(n) for n in zin.namelist()}

    pres = base_files["ppt/presentation.xml"].decode("utf-8")
    sld_ids = "".join([f'<p:sldId id="{255+i}" r:id="rId{10+i}"/>' for i in range(1, len(slides)+1)])
    pres = pres.replace("</p:sldMasterIdLst>", f"</p:sldMasterIdLst><p:sldIdLst>{sld_ids}</p:sldIdLst>")
    pres = pres.replace('<p:sldSz cx="9144000" cy="6858000" type="screen4x3"/>', '<p:sldSz cx="12192000" cy="6858000" type="wide"/>')

    ct = base_files["[Content_Types].xml"].decode("utf-8")
    if 'Extension="png"' not in ct:
        ct = ct.replace("</Types>", '<Default Extension="png" ContentType="image/png"/></Types>')
    slide_overrides = "".join([f'<Override PartName="/ppt/slides/slide{i}.xml" ContentType="application/vnd.openxmlformats-officedocument.presentationml.slide+xml"/>' for i in range(1, len(slides)+1)])
    ct = ct.replace("</Types>", slide_overrides + "</Types>")

    pres_rels = base_files["ppt/_rels/presentation.xml.rels"].decode("utf-8")
    new_rels = "".join([f'<Relationship Id="rId{10+i}" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/slide" Target="slides/slide{i}.xml"/>' for i in range(1, len(slides)+1)])
    pres_rels = pres_rels.replace("</Relationships>", new_rels + "</Relationships>")

    with zipfile.ZipFile(ppt, "w", zipfile.ZIP_DEFLATED) as z:
        for name, data in base_files.items():
            if name in {"[Content_Types].xml", "ppt/presentation.xml", "ppt/_rels/presentation.xml.rels"}:
                continue
            z.writestr(name, data)
        z.writestr("[Content_Types].xml", ct)
        z.writestr("ppt/presentation.xml", pres)
        z.writestr("ppt/_rels/presentation.xml.rels", pres_rels)
        for i, sx in enumerate(slides, start=1):
            z.writestr(f"ppt/slides/slide{i}.xml", sx)
            extra = rels[i-1].replace("</Relationships>", '<Relationship Id="rIdLayout" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/slideLayout" Target="../slideLayouts/slideLayout7.xml"/></Relationships>') if rels[i-1] else '<?xml version="1.0" encoding="UTF-8" standalone="yes"?><Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rIdLayout" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/slideLayout" Target="../slideLayouts/slideLayout7.xml"/></Relationships>'
            z.writestr(f"ppt/slides/_rels/slide{i}.xml.rels", extra)
        for path in paths.values():
            z.write(path, f"ppt/media/{path.name}")
    return ppt


def main():
    paths = {
        "arch": architecture_png(),
        "pipeline": pipeline_png(),
        "sequence": sequence_png(),
        "deployment": deployment_png(),
    }
    build_tech(paths)
    build_arch(paths)
    make_pptx(paths)
    for f in sorted(OUT.glob("*.*")):
        if f.suffix in [".docx", ".pptx"]:
            print(f.resolve())


if __name__ == "__main__":
    main()
