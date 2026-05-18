import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).parent / ".vendor"))

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image, ImageDraw, ImageFont


OUT = Path("交付物")
IMG = OUT / "common_agent_images"
OUT.mkdir(exist_ok=True)
IMG.mkdir(exist_ok=True)

FONT_MED = "/System/Library/Fonts/STHeiti Medium.ttc"
FONT_LIGHT = "/System/Library/Fonts/STHeiti Light.ttc"


def font(size, bold=False):
    return ImageFont.truetype(FONT_MED if bold else FONT_LIGHT, size)


def arrow(d, start, end, color="#64748B", width=4):
    import math
    d.line([start, end], fill=color, width=width)
    x1, y1 = start
    x2, y2 = end
    a = math.atan2(y2 - y1, x2 - x1)
    l = 15
    p1 = (x2 + l * math.cos(a + 2.45), y2 + l * math.sin(a + 2.45))
    p2 = (x2 + l * math.cos(a - 2.45), y2 + l * math.sin(a - 2.45))
    d.polygon([(x2, y2), p1, p2], fill=color)


def box(d, xy, title, lines=None, fill="#F8FAFC", outline="#334155", title_color="#111827", line_size=18):
    x1, y1, x2, y2 = xy
    d.rounded_rectangle(xy, radius=18, fill=fill, outline=outline, width=3)
    d.text((x1 + 18, y1 + 13), title, font=font(23, True), fill=title_color)
    y = y1 + 52
    for line in lines or []:
        d.text((x1 + 18, y), line, font=font(line_size), fill="#334155")
        y += line_size + 8


def save_framework_arch():
    img = Image.new("RGB", (1900, 1280), "#FFFFFF")
    d = ImageDraw.Draw(img)
    d.text((60, 42), "公共底层智能体框架总体架构", font=font(46, True), fill="#0F172A")
    d.text((60, 104), "用同一个 Agent OS 支撑“指标体系构建 Agent”和“离线 RAG / 字段填充 Agent”，后续新 Agent 通过模板与插件扩展。", font=font(24), fill="#475569")

    layers = [
        ("业务 Agent 应用层", ["指标体系构建 Agent", "RAG 问答/字段填充 Agent", "后续：合同/报表/数据治理 Agent"], "#DBEAFE", "#2563EB"),
        ("Agent Runtime 层", ["LangGraph 1.x 状态机", "Planner / Router / Executor", "Human-in-the-loop", "Checkpoint / Replay"], "#DCFCE7", "#16A34A"),
        ("能力插件层", ["Tool Registry", "MCP Server", "Skill Library", "Prompt/Schema Registry"], "#FEF3C7", "#D97706"),
        ("知识与数据服务层", ["文档解析 Pipeline", "Schema/字段字典", "混合检索", "业务数据查询"], "#F3E8FF", "#7C3AED"),
        ("模型网关层", ["LLM Router", "Embedding/Rerank", "CPU/API/GPU 模式", "JSON Schema 约束输出"], "#FFE4E6", "#E11D48"),
        ("基础设施与治理层", ["PostgreSQL + pgvector", "Redis/Queue", "审计脱敏", "可观测/评测/离线包"], "#E0F2FE", "#0284C7"),
    ]
    y = 170
    for i, (name, items, fill, outline) in enumerate(layers):
        d.rounded_rectangle((70, y, 1830, y + 135), radius=24, fill=fill, outline=outline, width=3)
        d.text((105, y + 43), name, font=font(31, True), fill="#111827")
        x = 410
        for item in items:
            w = 315 if len(items) == 4 else 420
            d.rounded_rectangle((x, y + 34, x + w, y + 94), radius=16, fill="#FFFFFF", outline=outline, width=2)
            d.text((x + 18, y + 52), item, font=font(20), fill="#1F2937")
            x += w + 28
        if i < len(layers) - 1:
            arrow(d, (950, y + 135), (950, y + 158), outline, 5)
        y += 158

    d.rounded_rectangle((70, 1172, 1830, 1232), radius=18, fill="#F8FAFC", outline="#CBD5E1", width=2)
    d.text((105, 1188), "设计原则：业务 Agent 薄、公共底座厚；确定性逻辑下沉到 Tool/Skill；模型只负责推理和生成；全链路可审计、可回放、可评测。", font=font(23), fill="#334155")
    p = IMG / "01_common_framework_arch.png"
    img.save(p)
    return p


def save_two_agent_mapping():
    img = Image.new("RGB", (1900, 1180), "#FFFFFF")
    d = ImageDraw.Draw(img)
    d.text((60, 42), "两个需求在公共框架上的落地关系", font=font(46, True), fill="#0F172A")
    d.text((60, 104), "差异在业务编排模板和工具组合，公共能力完全复用：模型网关、知识服务、工具注册、审计脱敏、可观测。", font=font(24), fill="#475569")

    d.rounded_rectangle((70, 170, 895, 950), radius=26, fill="#EFF6FF", outline="#2563EB", width=4)
    d.text((105, 205), "A. 指标体系构建 Agent", font=font(32, True), fill="#1D4ED8")
    metric_steps = [
        ("输入", ["表结构 JSON / Excel / Word", "业务 prompt / 数据域"]),
        ("Schema 瘦身", ["上千/上万表不拼大 JSON", "按库/表/域分批 Map"]),
        ("业务域分类", ["小模型路由", "字段字典 + RAG 辅助"]),
        ("指标推导", ["按业务域并行", "Tool: RAG / SchemaLookup"]),
        ("评估确认", ["可行性复评", "人工确认 / 版本化"]),
        ("输出", ["标准指标体系 JSON", "指标定义/口径/维度/SQL建议"]),
    ]
    y = 270
    for i, (t, lines) in enumerate(metric_steps):
        box(d, (125, y, 840, y + 88), t, lines, "#FFFFFF", "#60A5FA", line_size=15)
        if i < len(metric_steps) - 1:
            arrow(d, (482, y + 88), (482, y + 108), "#2563EB", 4)
        y += 110

    d.rounded_rectangle((1005, 170, 1830, 950), radius=26, fill="#F0FDF4", outline="#16A34A", width=4)
    d.text((1040, 205), "B. 离线 RAG / 字段填充 Agent", font=font(32, True), fill="#15803D")
    rag_steps = [
        ("输入", ["Excel/CSV/Word/PPT/PDF", "TXT/Markdown/HTML"]),
        ("文件解析", ["按类型路由解析器", "Excel 结构化解析"]),
        ("知识写入", ["行级 chunk + row_json", "pgvector + 普通索引"]),
        ("查询理解", ["字段名/字段值抽取", "别名映射"]),
        ("检索路由", ["精确查询优先", "向量召回补充"]),
        ("输出", ["字段值填充 / 问答", "引用、置信度、审计"]),
    ]
    y = 270
    for i, (t, lines) in enumerate(rag_steps):
        box(d, (1060, y, 1775, y + 88), t, lines, "#FFFFFF", "#4ADE80", line_size=15)
        if i < len(rag_steps) - 1:
            arrow(d, (1417, y + 88), (1417, y + 108), "#16A34A", 4)
        y += 110

    d.rounded_rectangle((220, 1000, 1680, 1130), radius=24, fill="#FEF3C7", outline="#D97706", width=3)
    d.text((250, 1024), "公共复用底座", font=font(30, True), fill="#92400E")
    d.text((250, 1070), "Agent Runtime / Tool Registry / MCP / Skill Library / RAG Service / Model Gateway / Audit / Evaluation / Offline Deployment", font=font(24), fill="#78350F")
    arrow(d, (482, 950), (770, 1000), "#D97706", 5)
    arrow(d, (1417, 950), (1130, 1000), "#D97706", 5)

    p = IMG / "02_two_agent_mapping.png"
    img.save(p)
    return p


def save_agent_runtime_flow():
    img = Image.new("RGB", (1900, 1100), "#FFFFFF")
    d = ImageDraw.Draw(img)
    d.text((60, 42), "公共 Agent Runtime 状态机", font=font(46, True), fill="#0F172A")
    d.text((60, 104), "所有业务 Agent 都按统一生命周期运行：输入规范化、计划、工具执行、校验、人工确认、输出、审计回放。", font=font(24), fill="#475569")
    nodes = [
        ("Input Adapter", "文件/JSON/Prompt/API\n统一转 TaskContext", 130, 230, "#DBEAFE", "#2563EB"),
        ("Policy Guard", "鉴权/脱敏/数据域\nPrompt 注入防护", 480, 230, "#FFE4E6", "#E11D48"),
        ("Planner", "选择 Agent 模板\n拆解步骤/并行度", 830, 230, "#DCFCE7", "#16A34A"),
        ("Tool Executor", "调用 RAG/DB/Schema\nMCP/Skill/内部服务", 1180, 230, "#FEF3C7", "#D97706"),
        ("Model Gateway", "小模型/大模型/Embedding\nJSON Schema 输出", 1530, 230, "#F3E8FF", "#7C3AED"),
        ("Validator", "Pydantic 校验\n事实/引用/权限检查", 1180, 565, "#E0F2FE", "#0284C7"),
        ("HITL", "人工确认\n修正/驳回/继续", 830, 565, "#FEF3C7", "#D97706"),
        ("Answer Builder", "结构化 JSON\n答案/引用/置信度", 480, 565, "#DCFCE7", "#16A34A"),
        ("Audit & Replay", "Trace/日志/评测样本\n可回放可复盘", 130, 565, "#F8FAFC", "#334155"),
    ]
    centers = {}
    for title, lines, x, y, fill, outline in nodes:
        box(d, (x, y, x + 245, y + 150), title, lines.split("\n"), fill, outline)
        centers[title] = (x + 122, y + 75)
    sequence = ["Input Adapter", "Policy Guard", "Planner", "Tool Executor", "Model Gateway", "Validator", "HITL", "Answer Builder", "Audit & Replay"]
    for a, b in zip(sequence, sequence[1:]):
        arrow(d, centers[a], centers[b], "#475569", 4)
    arrow(d, centers["Validator"], centers["Tool Executor"], "#DC2626", 3)
    d.text((1280, 505), "校验失败：回到工具/模型重试", font=font(19), fill="#DC2626")
    d.rounded_rectangle((130, 860, 1775, 1015), radius=24, fill="#F8FAFC", outline="#CBD5E1", width=3)
    d.text((170, 890), "Runtime 统一状态对象 TaskContext", font=font(30, True), fill="#111827")
    d.text((170, 940), "task_id、agent_type、input_payload、security_context、plan、tool_calls、retrieval_context、model_outputs、validation_errors、human_feedback、final_result", font=font(22), fill="#334155")
    p = IMG / "03_agent_runtime_flow.png"
    img.save(p)
    return p


def save_extension_lifecycle():
    img = Image.new("RGB", (1900, 1040), "#FFFFFF")
    d = ImageDraw.Draw(img)
    d.text((60, 42), "新业务 Agent 扩展生命周期", font=font(46, True), fill="#0F172A")
    d.text((60, 104), "后续扩展不是复制工程，而是注册模板、工具、Schema、评测集与权限策略。", font=font(24), fill="#475569")
    steps = [
        ("1. 定义业务模板", ["Agent 类型", "状态机 DAG", "输入/输出 Schema"], "#DBEAFE", "#2563EB"),
        ("2. 注册工具", ["Tool manifest", "MCP Server", "Skill 流程"], "#FEF3C7", "#D97706"),
        ("3. 绑定知识域", ["RAG collection", "字段字典", "权限范围"], "#F3E8FF", "#7C3AED"),
        ("4. 配置模型策略", ["路由规则", "Prompt 模板", "结构化输出"], "#DCFCE7", "#16A34A"),
        ("5. 建立评测集", ["黄金样本", "回归指标", "人工反馈"], "#FFE4E6", "#E11D48"),
        ("6. 发布与灰度", ["版本号", "审计策略", "回滚开关"], "#E0F2FE", "#0284C7"),
    ]
    x, y = 90, 220
    centers = []
    for i, (title, lines, fill, outline) in enumerate(steps):
        box(d, (x, y, x + 500, y + 185), title, lines, fill, outline)
        centers.append((x + 250, y + 92))
        if i % 3 != 2:
            arrow(d, (x + 500, y + 92), (x + 575, y + 92), outline, 5)
        x += 590
        if i == 2:
            x, y = 90, 535
    arrow(d, (1680, 312), (1680, 628), "#64748B", 5)
    arrow(d, (590, 628), (665, 628), "#64748B", 5)
    d.rounded_rectangle((90, 830, 1810, 970), radius=24, fill="#F8FAFC", outline="#CBD5E1", width=3)
    d.text((130, 860), "扩展边界", font=font(30, True), fill="#111827")
    d.text((130, 910), "新增业务时只新增 Agent 模板、Tool/Skill、Schema、评测样本；不改 Agent Runtime、模型网关、知识服务、审计与部署底座。", font=font(23), fill="#334155")
    p = IMG / "04_extension_lifecycle.png"
    img.save(p)
    return p


def save_deployment():
    img = Image.new("RGB", (1900, 1080), "#FFFFFF")
    d = ImageDraw.Draw(img)
    d.text((60, 42), "公共智能体平台离线部署架构", font=font(46, True), fill="#0F172A")
    d.text((60, 104), "单机虚机起步，CPU 功能验证；确认模型后迁移到 GPU 推理服务，业务 Agent 无需改造。", font=font(24), fill="#475569")
    d.rounded_rectangle((70, 170, 1830, 925), radius=28, fill="#F8FAFC", outline="#334155", width=4)
    comps = [
        ((120, 260, 445, 410), "Gateway", ["Nginx/APISIX", "鉴权/限流/灰度"], "#DBEAFE", "#2563EB"),
        ((540, 260, 875, 410), "Agent API", ["FastAPI", "Agent Template Router"], "#DCFCE7", "#16A34A"),
        ((970, 260, 1305, 410), "Agent Worker", ["LangGraph Runtime", "任务队列/断点恢复"], "#DCFCE7", "#16A34A"),
        ((1400, 260, 1730, 410), "Model Gateway", ["CPU/API/GPU", "统一 OpenAI 接口"], "#F3E8FF", "#7C3AED"),
        ((120, 545, 445, 710), "Knowledge Service", ["文档解析/分片", "RAG/字段字典"], "#FEF3C7", "#D97706"),
        ((540, 545, 875, 710), "PostgreSQL 16", ["pgvector/audit", "excel_rows/metrics"], "#FFE4E6", "#E11D48"),
        ((970, 545, 1305, 710), "Redis / Queue", ["缓存/队列", "限流计数"], "#FFEDD5", "#EA580C"),
        ((1400, 545, 1730, 710), "Observability", ["Prometheus/Grafana", "Phoenix/Loki"], "#E0F2FE", "#0284C7"),
    ]
    for xy, title, lines, fill, outline in comps:
        box(d, xy, title, lines, fill, outline)
    for s, e in [((445, 335), (540, 335)), ((875, 335), (970, 335)), ((1305, 335), (1400, 335)), ((1137, 410), (1137, 545)), ((707, 410), (707, 545)), ((282, 410), (282, 545))]:
        arrow(d, s, e, "#475569", 4)
    d.text((120, 805), "离线包：Docker 镜像 tar、Python wheelhouse、模型权重、初始化 SQL、Agent 模板包、Tool/MCP manifest、评测样本、验收脚本。", font=font(24), fill="#334155")
    d.text((120, 850), "最低配置：CPU 功能验证 16C/64GB/500GB；增强验证 32C/128GB/1TB；GPU 生产按模型压测选择 L20 48GB / A100 40GB 起。", font=font(24), fill="#334155")
    p = IMG / "05_deployment.png"
    img.save(p)
    return p


def shade(cell, color):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), color)
    tc_pr.append(shd)


def setup_doc(doc):
    sec = doc.sections[0]
    sec.top_margin = Inches(0.65)
    sec.bottom_margin = Inches(0.65)
    sec.left_margin = Inches(0.72)
    sec.right_margin = Inches(0.72)
    for s in ["Normal", "Heading 1", "Heading 2", "Heading 3"]:
        doc.styles[s].font.name = "Arial Unicode MS"
    doc.styles["Normal"].font.size = Pt(10.5)
    doc.styles["Heading 1"].font.size = Pt(18)
    doc.styles["Heading 2"].font.size = Pt(14)
    doc.styles["Heading 3"].font.size = Pt(12)


def add_title(doc, title, sub):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(title)
    r.bold = True
    r.font.size = Pt(22)
    r.font.color.rgb = RGBColor(15, 23, 42)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(sub)
    r.font.size = Pt(11)
    r.font.color.rgb = RGBColor(71, 85, 105)
    doc.add_paragraph("")


def bullets(doc, items):
    for item in items:
        doc.add_paragraph(item, style="List Bullet")


def table(doc, headers, rows):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(headers):
        c = t.rows[0].cells[i]
        c.text = h
        c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        shade(c, "E2E8F0")
        for p in c.paragraphs:
            for r in p.runs:
                r.bold = True
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = str(v)
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
    doc.add_paragraph("")


def image(doc, path, caption):
    doc.add_picture(str(path), width=Inches(6.85))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = doc.add_paragraph(caption)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for r in p.runs:
        r.font.size = Pt(9)
        r.font.color.rgb = RGBColor(100, 116, 139)


def build_arch_doc(paths):
    doc = Document()
    setup_doc(doc)
    add_title(doc, "公共底层智能体框架 · 架构设计文档", "面向指标体系构建 Agent 与离线 RAG/字段填充 Agent 的统一架构")

    doc.add_heading("1. 修正后的需求理解", 1)
    bullets(doc, [
        "本次目标不是只为某一个 RAG 场景画架构，而是设计一个公共底层智能体框架，用同一套平台能力承载多个业务 Agent。",
        "需求一：指标体系构建 Agent。输入表结构 JSON、Excel/Word 文档或业务 prompt，解决大 JSON 拼接瓶颈，完成表结构理解、业务域分类、RAG 辅助、指标推导、可行性复评、人工确认和标准指标体系 JSON 输出。",
        "需求二：离线 RAG/字段填充 Agent。输入 Excel/CSV、Word、PPT、PDF、TXT 等文件，完成离线解析、行级 chunk、pgvector 写入、字段别名映射、精确查询、语义召回、字段填充、问答和审计反馈。",
        "两个需求的差异是业务流程与工具组合；共同能力是 Agent Runtime、Tool/MCP/Skill、模型路由、知识服务、审计脱敏、可观测、离线部署和评测闭环。",
    ])

    doc.add_heading("2. 设计原则", 1)
    table(doc, ["原则", "说明"], [
        ("业务 Agent 薄，公共底座厚", "业务层只定义模板、状态机、输入输出 Schema 与工具编排，不重复建设模型、检索、审计、部署能力。"),
        ("确定性逻辑工具化", "文件解析、SQL 查询、字段填充、JSON 校验、权限过滤等确定性逻辑下沉到 Tool/Skill，避免让 LLM 处理脏活。"),
        ("模型网关统一", "CPU 验证、云 API 选型、GPU 自部署通过同一 LLM Router 暴露统一接口，上层 Agent 不感知模型位置。"),
        ("知识服务标准化", "RAG、字段字典、Schema Registry、Excel row_json、指标知识库统一成为 Knowledge Service。"),
        ("全链路可审计可回放", "每个 request_id 串联输入、计划、工具调用、检索结果、模型输出、人工反馈、最终结果。"),
        ("扩展靠注册，不靠复制", "新增 Agent 通过注册 Agent 模板、工具 manifest、Schema、Prompt 和评测集完成。"),
    ])

    doc.add_heading("3. 总体架构", 1)
    image(doc, paths["arch"], "图 1：公共底层智能体框架总体架构")
    doc.add_paragraph("架构分为六层：业务 Agent 应用层、Agent Runtime 层、能力插件层、知识与数据服务层、模型网关层、基础设施与治理层。指标体系构建和离线 RAG/字段填充只是应用层的两个 Agent 模板，公共能力集中在下层复用。")

    doc.add_heading("4. 两个需求如何落到同一个框架", 1)
    image(doc, paths["mapping"], "图 2：两个业务 Agent 在公共框架上的落地关系")
    table(doc, ["能力域", "指标体系构建 Agent", "离线 RAG/字段填充 Agent", "公共底座复用点"], [
        ("输入适配", "表结构 JSON、Excel、Word、业务 prompt", "Excel/CSV、Word、PPT、PDF、TXT", "Input Adapter + DocumentModel + TaskContext"),
        ("大输入处理", "Schema 瘦身、分批 Map、按业务域 Reduce", "行级 chunk、metadata、row_json", "Chunking Policy + Schema Registry"),
        ("知识使用", "指标知识库、行业口径、表字段语义", "文档知识库、字段字典、Excel 行数据", "Knowledge Service + RAG Tool"),
        ("工具调用", "SchemaLookup、MetricRAG、SQLAdvisor、FeasibilityCheck", "ExactRowQuery、VectorSearch、FieldAlias、AnswerBuilder", "Tool Registry / MCP / Skill Library"),
        ("输出约束", "指标体系 JSON、口径、维度、SQL 建议", "字段值、整行记录、问答答案、引用", "Pydantic / JSON Schema / Validator"),
        ("治理", "指标版本、人工确认、评测样本", "审计、脱敏、错误样本回流", "Audit、HITL、Evaluation"),
    ])

    doc.add_heading("5. 公共 Agent Runtime 设计", 1)
    image(doc, paths["runtime"], "图 3：公共 Agent Runtime 状态机")
    bullets(doc, [
        "Input Adapter：把文件、JSON、Prompt、API 请求统一转换为 TaskContext。",
        "Policy Guard：做鉴权、数据域过滤、脱敏、Prompt 注入检测、工具权限控制。",
        "Planner：根据 agent_type 选择业务模板，生成步骤计划、并行度、工具调用策略。",
        "Tool Executor：执行 RAG 检索、字段查询、Schema 查询、指标推导辅助、审计写入等工具。",
        "Model Gateway：统一调用小模型、主模型、Embedding、Rerank；支持 CPU/API/GPU 三模式。",
        "Validator：对模型输出做 JSON Schema/Pydantic 校验、引用校验、权限校验与事实一致性检查。",
        "HITL：在指标确认、低置信度字段填充、冲突口径等节点引入人工确认。",
        "Audit & Replay：保存完整 Trace，支持问题复现、回归评测和版本对比。",
    ])

    doc.add_heading("6. Tool / MCP / Skill 分层", 1)
    table(doc, ["层级", "定位", "示例", "设计边界"], [
        ("Tool", "单个确定性能力，输入输出强 Schema", "pgvector_search、exact_row_query、schema_lookup、metric_rag_search、audit_write", "工具不做业务编排，只完成一个清晰动作。"),
        ("MCP Server", "跨 Agent 复用的外部能力封装", "database-mcp、file-mcp、knowledge-mcp、metrics-mcp", "适合把数据库、文件、内部服务标准化暴露。"),
        ("Skill", "可复用的专家流程或规则包", "excel_row_chunking、field_alias_learning、metric_feasibility_check、rag_quality_check", "适合沉淀业务规则，减少 Prompt 依赖。"),
        ("Agent Template", "业务级流程编排", "metric_system_builder、offline_rag_field_filler", "只组合 Tool/MCP/Skill，不内嵌底层实现。"),
    ])

    doc.add_heading("7. 扩展机制", 1)
    image(doc, paths["extension"], "图 4：新业务 Agent 扩展生命周期")
    bullets(doc, [
        "新增业务 Agent 时，不新建一套工程，只新增 agent.yaml、state_schema.py、tools.yaml、prompt 模板、输出 JSON Schema、评测集和权限策略。",
        "Agent Runtime 根据 agent_type 动态加载模板；Tool Registry 根据 manifest 控制工具可见性、参数 Schema、权限和审计级别。",
        "知识服务按 collection/domain/source_id 做隔离，避免不同 Agent 的知识域混用。",
        "模型策略按任务级别分流：路由/分类用小模型，复杂推理用主模型，检索用 embedding/rerank，字段填充优先不用 LLM。",
    ])

    doc.add_heading("8. 部署架构", 1)
    image(doc, paths["deployment"], "图 5：公共智能体平台离线部署架构")
    bullets(doc, [
        "单机虚机起步：Gateway、Agent API、Agent Worker、Knowledge Service、Model Gateway、PostgreSQL、Redis、可观测组件部署在同一节点。",
        "CPU 验证阶段只验证流程、Schema、工具、RAG 和审计闭环；7B GGUF 模型可用于低速验证，不作为性能基准。",
        "模型确认后再采购 GPU，Model Gateway 从 llama.cpp/API 模式切换到 vLLM 模式，业务 Agent 无需改造。",
        "内网离线交付包含 Docker 镜像 tar、Python wheelhouse、模型权重、初始化 SQL、Agent 模板包、Tool/MCP manifest、评测样本和验收脚本。",
    ])

    doc.add_heading("9. 数据与状态模型", 1)
    table(doc, ["对象", "关键字段", "说明"], [
        ("agent_templates", "agent_type, version, graph_config, input_schema, output_schema", "业务 Agent 模板注册表。"),
        ("agent_runs", "run_id, agent_type, status, user_id, task_context, final_result", "一次 Agent 执行的主记录。"),
        ("tool_calls", "run_id, tool_name, args, result_ref, latency, status", "工具调用明细，用于审计和回放。"),
        ("knowledge_collections", "collection_id, domain, source_type, acl_policy", "知识域管理，支持指标知识库和文档知识库隔离。"),
        ("rag_chunks", "chunk_id, collection_id, chunk_text, embedding, metadata", "向量检索主表。"),
        ("excel_rows", "row_id, source_id, sheet_name, row_no, row_json, indexed_fields", "Excel 字段填充主表。"),
        ("metric_artifacts", "artifact_id, run_id, domain, metric_json, review_status", "指标体系构建结果与版本。"),
        ("evaluation_cases", "case_id, agent_type, input, expected, scoring_rule", "回归评测样本。"),
    ])

    doc.save(OUT / "公共底层智能体框架_架构设计文档.docx")


def build_selection_doc(paths):
    doc = Document()
    setup_doc(doc)
    add_title(doc, "公共底层智能体框架 · 技术选型方案", "支撑指标体系构建与离线 RAG/字段填充两个业务需求的统一组件选型")

    doc.add_heading("1. 选型结论", 1)
    table(doc, ["层级", "推荐组件", "版本建议", "理由"], [
        ("开发语言", "Python + Java", "Python 3.12.x；Java 17 / Spring Boot 3.3.x", "Python 承载 AI 底座；Java 承载企业系统接入、权限和管理后台。"),
        ("Agent Runtime", "LangGraph", "1.2.0", "1.x 生产稳定，适合状态机、多步骤、可回放 Agent。"),
        ("LLM 抽象", "LangChain", "1.3.1", "模型/工具/消息抽象成熟，配合 LangGraph 使用；只作为适配层，避免业务重度耦合。"),
        ("API 服务", "FastAPI + Uvicorn", "FastAPI 0.136.1；Uvicorn 0.34.x+", "异步 API、SSE、Pydantic 生态好，适合 Agent 服务。"),
        ("数据校验", "Pydantic", "2.12.x+", "统一定义 TaskContext、Tool 参数、Agent 输出 JSON Schema。"),
        ("任务队列", "Celery / RQ + Redis", "Celery 5.5.x+；Redis 7.4.x", "离线解析、embedding、批量任务、重试与缓存。"),
        ("数据库", "PostgreSQL", "16.x / 17.x", "统一承载业务表、审计、JSONB、全文索引和向量扩展。"),
        ("向量扩展", "pgvector", "0.8.2+", "包含 HNSW 能力，并修复并行 HNSW 构建相关安全问题。"),
        ("文档解析", "pandas/openpyxl/python-docx/python-pptx/PyMuPDF/unstructured", "pandas 2.2.x；openpyxl 3.1.x；PyMuPDF 1.24.x+", "覆盖结构化 Excel 与非结构化文档解析。"),
        ("Embedding/Rerank", "bge-m3 + bge-reranker-v2-m3", "离线模型文件固定版本", "中文、多语言和企业文档检索效果稳定。"),
        ("CPU 推理", "llama.cpp / llama-cpp-python", "0.3.x+", "前期 CPU 验证功能链路。"),
        ("GPU 推理", "vLLM", "0.8.x+", "确认模型后用于 GPU 高吞吐推理，OpenAI 兼容接口。"),
        ("审计脱敏", "Presidio + 自定义规则", "Presidio 2.2.x+", "满足手机号、身份证、邮箱、合同号、客户名等脱敏需求。"),
        ("可观测", "OpenTelemetry + Prometheus + Grafana + Phoenix", "OTel 1.30.x+", "指标、日志、Trace、LLM 调用链路和评测闭环。"),
    ])

    doc.add_heading("2. 为什么是公共底座，而不是两个独立系统", 1)
    table(doc, ["对比项", "两个独立系统", "公共智能体框架"], [
        ("建设成本", "解析、检索、模型、审计各做一套", "公共能力一次建设，多 Agent 复用"),
        ("扩展性", "新增需求继续复制工程", "新增 Agent 模板 + Tool/Skill 注册"),
        ("治理", "日志、权限、评测分散", "统一审计、统一权限、统一评测"),
        ("模型切换", "每个系统各自适配", "Model Gateway 统一切换 CPU/API/GPU"),
        ("知识复用", "指标知识、文档知识割裂", "Knowledge Service 统一 collection/domain 管理"),
        ("风险", "烟囱化，长期维护成本高", "底座复杂度更高，但治理清晰、长期收益大"),
    ])

    doc.add_heading("3. Python 与 Java 技术栈对比", 1)
    table(doc, ["维度", "Java 方案", "Python 方案", "推荐"], [
        ("公司基础", "后端团队熟悉，工程治理强", "AI 工程规范需要建立", "Java 保留接入层和管理后台"),
        ("Agent/RAG 生态", "Spring AI、LangChain4j 可用，但新能力跟进较慢", "LangGraph/LangChain/vLLM/Transformers/Embedding 生态完整", "公共 Agent 底座选 Python"),
        ("文档解析", "POI/Tika 稳定但组合复杂", "pandas/openpyxl/PyMuPDF/unstructured 更灵活", "解析与 RAG Pipeline 选 Python"),
        ("模型推理", "通常调用外部模型服务", "CPU/GPU 推理生态成熟", "Model Gateway 选 Python"),
        ("企业集成", "权限、组织、审批、门户集成优势强", "需额外开发", "企业系统集成仍用 Java"),
        ("最终建议", "不作为 AI 核心底座主语言", "作为智能体底座主语言", "Java + Python 边界清晰解耦"),
    ])

    doc.add_heading("4. 核心组件如何满足两个需求", 1)
    table(doc, ["组件", "支撑指标体系构建", "支撑离线 RAG/字段填充"], [
        ("LangGraph", "业务域分类、并行指标推导、可行性复评、人工确认", "离线写入、查询理解、检索路由、答案生成、审计反馈"),
        ("Tool Registry", "SchemaLookup、MetricRAG、SQLAdvisor、FeasibilityCheck", "ExactRowQuery、VectorSearch、FieldAlias、AnswerBuilder"),
        ("Knowledge Service", "指标口径库、表字段语义库、行业参考库", "文档知识库、Excel row_json、字段字典"),
        ("Model Gateway", "小模型分类、主模型推导、结构化 JSON 输出", "意图识别、问答生成、Embedding/Rerank"),
        ("PostgreSQL + pgvector", "存储 schema、指标产物、指标知识 chunk", "存储 rag_chunks、excel_rows、字段索引、审计"),
        ("HITL", "指标口径确认、可行性确认、版本发布", "低置信度字段填充确认、错误反馈"),
        ("Evaluation", "指标 JSON 完整性、口径一致性、SQL 可执行性", "字段填充准确率、召回率、引用正确性"),
    ])

    doc.add_heading("5. 模型策略", 1)
    table(doc, ["阶段", "模型/方式", "目标", "说明"], [
        ("CPU 功能验证", "Qwen2.5-7B-Instruct GGUF Q4_K_M + llama.cpp", "验证 Agent、Tool、RAG、Schema、审计闭环", "不以延迟作为性能指标。"),
        ("API 选型验证", "阿里百炼/火山方舟等同类 API 模型", "比较 7B/14B/32B 在指标推导与查询问答上的效果", "确认效果后再决定 GPU。"),
        ("GPU 自部署", "Qwen2.5-14B/32B-Instruct-AWQ 或 Qwen3-30B-A3B + vLLM", "支撑生产吞吐和成本控制", "以压测结果确定显卡型号。"),
        ("Embedding/Rerank", "bge-m3 + bge-reranker-v2-m3", "保障 RAG 召回与排序", "离线部署，版本固化。"),
    ])

    doc.add_heading("6. 最低环境配置", 1)
    table(doc, ["环境", "配置", "可验证内容", "注意事项"], [
        ("CPU 最小验证", "16C / 64GB RAM / 500GB SSD", "API、Agent Runtime、Tool、RAG、审计、7B 慢速推理", "只验证功能，不承诺响应时间。"),
        ("CPU 推荐验证", "32C / 128GB RAM / 1TB SSD", "批量文档导入、几万字段索引、100 QPS 非模型接口压测", "适合内网 PoC。"),
        ("GPU 生产起步", "32C / 128GB RAM / 2TB SSD + L20 48GB 或 A100 40GB", "14B/32B 量化模型推理、Embedding/Rerank 服务", "最终以模型压测确定。"),
        ("离线包", "Docker tar + wheelhouse + 模型权重 + SQL + Agent 模板", "无外网安装部署", "全部组件需锁版本和校验 hash。"),
    ])

    doc.add_heading("7. 版本与安全说明", 1)
    bullets(doc, [
        "LangChain 采用 1.3.1，LangGraph 采用 1.2.0，均满足 1.0 以上要求。",
        "pgvector 采用 0.8.2+，原因是该版本修复了并行 HNSW 索引构建相关安全问题。",
        "离线环境必须建立私有 wheelhouse 和镜像仓库，不能部署时临时联网安装。",
        "所有 Tool 入参必须做 Schema 校验，禁止用户输入直接拼接 SQL、metadata filter 或文件路径。",
        "LLM 输出一律视为不可信输入，必须经过 Validator 校验后才能入库或展示。",
    ])

    doc.add_heading("8. 参考来源", 1)
    bullets(doc, [
        "LangChain PyPI：langchain 1.3.1，2026-05-15 发布。",
        "LangGraph PyPI：langgraph 1.2.0，2026-05-12 发布。",
        "FastAPI PyPI：fastapi 0.136.1，2026-04-23 发布。",
        "PostgreSQL 官方新闻：pgvector 0.8.2 于 2026-02-26 发布并修复 CVE-2026-3172。",
    ])

    doc.save(OUT / "公共底层智能体框架_技术选型方案.docx")


def main():
    paths = {
        "arch": save_framework_arch(),
        "mapping": save_two_agent_mapping(),
        "runtime": save_agent_runtime_flow(),
        "extension": save_extension_lifecycle(),
        "deployment": save_deployment(),
    }
    build_arch_doc(paths)
    build_selection_doc(paths)
    for p in sorted(OUT.glob("公共底层智能体框架_*.docx")):
        print(p.resolve())


if __name__ == "__main__":
    main()
