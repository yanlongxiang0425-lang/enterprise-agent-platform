from pathlib import Path
import sys

sys.path.insert(0, str(Path.cwd() / ".vendor"))
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn


OUT = Path("交付物")
ARCH_DOC = OUT / "公共底层智能体平台技术架构设计文档_物料样例增强版.docx"
SELECT_DOC = OUT / "公共底层智能体平台技术选型方案_物料样例增强版.docx"


def setup(doc: Document):
    sec = doc.sections[0]
    sec.top_margin = Inches(0.75)
    sec.bottom_margin = Inches(0.75)
    sec.left_margin = Inches(0.75)
    sec.right_margin = Inches(0.75)
    styles = doc.styles
    styles["Normal"].font.name = "Arial"
    styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    styles["Normal"].font.size = Pt(10.5)
    for name, size, color in [("Heading 1", 16, "2E74B5"), ("Heading 2", 13, "1F4D78")]:
        styles[name].font.name = "Arial"
        styles[name]._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
        styles[name].font.size = Pt(size)
        styles[name].font.color.rgb = RGBColor.from_string(color)


def title(doc, text, sub):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    r.bold = True
    r.font.size = Pt(20)
    r.font.color.rgb = RGBColor(15, 23, 42)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(sub)
    r.font.size = Pt(10)
    r.font.color.rgb = RGBColor(100, 116, 139)


def table(doc, headers, rows):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    for i, h in enumerate(headers):
        c = t.rows[0].cells[i]
        c.text = h
        for p in c.paragraphs:
            for r in p.runs:
                r.bold = True
                r.font.size = Pt(9)
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = str(v)
            for p in cells[i].paragraphs:
                for r in p.runs:
                    r.font.size = Pt(9)
    doc.add_paragraph()


def bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.add_run(item)


def build_selection():
    doc = Document()
    setup(doc)
    title(doc, "公共底层智能体平台技术选型方案", "物料分类补全真实样例增强版")
    doc.add_heading("1. 样例数据理解", 1)
    doc.add_paragraph("本次补充分析了两组真实测试数据，技术选型需要围绕“Excel结构化处理 + 字段映射 + 物料知识库精确补全 + RAG/模型辅助判断”展开。")
    table(doc, ["目录", "文件", "作用", "系统含义"], [
        ["试点物料 AI", "IOT PLM TV整机-TV成品-TV整机  100201.xls", "知识库导入数据", "PLM物料明细，包含WTPart/NonRDPart Sheet、编号、旧物料号、短描述、长描述、分类、尺寸、品牌等72列左右字段"],
        ["试点物料 AI", "10整机-TV成品-TV整机 目标模板.xlsx", "目标模板", "定义主数据标准、值集收集、清洗目标模板，约束需要补全的字段、拼接规则、必填规则和值集规则"],
        ["试点物料 AI", "10整机-TV成品-TV整机-数据清洗试点_何海珍 20260509.xlsx", "补全样例输出", "目标模板被补全后的参考形态，用于反推输出结构和验收字段"],
        ["AI测试文件", "第1-5步*.xls", "字段映射与清洗步骤", "旧字段标准 + 新字段标准生成映射表；旧料号信息 + 映射表生成清洗结果表"],
    ])
    doc.add_heading("2. 技术选型结论", 1)
    table(doc, ["组件", "建议版本", "作用", "选择理由"], [
        ["Python", "3.11.x", "agent-api统一AI服务", "Excel处理、RAG、LangGraph、模型调用生态更成熟；Java后台继续保留管理能力"],
        ["FastAPI", "0.115+", "REST/SSE接口", "轻量、异步、适合文件上传、任务进度和结果下载"],
        ["LangGraph", "1.0+", "Agent状态机编排", "将Excel解析、字段识别、候选召回、分类判断、属性补全、校验导出拆成可恢复节点"],
        ["LangChain", "1.0+", "模型与RAG工具抽象", "用于模型调用、Embedding/Rerank适配，避免绑定单一厂商"],
        ["pandas", "2.0.3", "Excel数据处理", "适合字段标准、映射表、旧料号信息和补全结果的表格变换"],
        ["openpyxl", "3.1.5", ".xlsx读写", "用于目标模板、补全结果输出、异常Sheet和统计Sheet"],
        ["xlrd", "2.0.1", ".xls读取", "真实样例大量为WPS生成的xls老格式，需要支持读取"],
        ["PostgreSQL + pgvector", "PostgreSQL 16 + pgvector 0.7+", "物料结构化库和辅助向量库", "主路径走结构化精确匹配；RAG用于规则解释、语义辅助、低置信度兜底"],
        ["Redis", "7.x", "任务队列、缓存、checkpoint辅助", "支持大Excel异步处理、断点续跑、任务状态查询"],
        ["Model Gateway", "agent-api进程内模块", "模型路由、限流、降级", "前期不拆独立服务，统一收敛在agent-api，降低单机离线部署复杂度"],
    ])
    doc.add_heading("3. 选型约束", 1)
    bullets(doc, [
        "物料补全主链路不应设计为纯RAG问答；应以结构化物料库、字段映射、规则拼接和精确查询为主。",
        "RAG和模型只处理字段语义匹配、分类规则解释、候选冲突、低置信度判断等辅助场景。",
        "前期单机VM部署，AI能力合并在agent-api进程内，后续QPS或模型推理压力上来后再拆模型服务。",
        "必须支持.xls和.xlsx两类Excel输入，保留原行号、原字段、异常原因、置信度和来源。",
    ])
    doc.save(SELECT_DOC)


def build_arch():
    doc = Document()
    setup(doc)
    title(doc, "公共底层智能体平台技术架构设计文档", "物料分类补全真实样例增强版")
    doc.add_heading("1. 真实需求链路复述", 1)
    doc.add_paragraph("物料分类补全Agent的目标是：导入PLM物料知识库，读取目标模板要求，基于物料编号、旧物料号、字段映射、分类规则和拼接规则，自动生成补全后的主数据标准/清洗结果Excel。")
    table(doc, ["链路", "输入", "处理", "输出"], [
        ["试点物料三文件", "PLM知识库xls + 目标模板xlsx", "结构化导入、字段映射、属性补全、异常识别", "补全结果Excel、异常Sheet、统计Sheet"],
        ["AI测试五步", "旧字段标准 + 新字段标准 + 旧料号信息", "生成新旧字段映射表，再按映射表抽取旧料号属性值", "新旧字段映射表、清洗结果表"],
    ])
    doc.add_heading("2. agent-api内部模块", 1)
    table(doc, ["模块", "职责", "对应样例"], [
        ["Excel Parser", "识别Sheet、表头、数据行、空列、WPS .xls/.xlsx差异", "所有8个Excel文件"],
        ["Knowledge Loader", "导入PLM WTPart/NonRDPart物料明细，建立编号/旧物料号索引", "IOT PLM TV整机...xls"],
        ["Template Analyzer", "读取主数据标准、值集收集、清洗目标模板，抽取目标字段和补全规则", "目标模板xlsx"],
        ["Field Mapping Engine", "旧字段标准、新字段标准、映射表、字段别名和单位映射", "AI测试文件第1-3步"],
        ["Classification/Completion Node", "规则优先、结构化匹配、模型辅助，生成分类和属性补全结果", "试点物料补全"],
        ["Validation Node", "检查未映射字段、缺失值、多命中、低置信度、覆盖冲突", "异常Sheet"],
        ["Export Tool", "输出补全结果、异常明细、差异报告、统计Sheet", "样例输出文件"],
    ])
    doc.add_heading("3. LangGraph节点链路", 1)
    bullets(doc, [
        "parse_excel：解析知识库、模板、待补全文件，保留sheet、行号、原字段。",
        "detect_schema：识别PLM知识库字段、目标模板字段、新旧字段标准字段。",
        "load_knowledge：将PLM明细结构化入库，建立编号、旧物料号、分类、描述字段索引。",
        "build_mapping：生成或读取新旧字段映射表，补充字段别名和单位规则。",
        "retrieve_candidate：按物料编号/旧物料号精确匹配，按名称/描述/分类做候选召回。",
        "complete_attributes：按目标模板字段输出补全值、来源字段、置信度、补全策略。",
        "validate_export：生成异常明细、统计结果和最终Excel。",
    ])
    doc.add_heading("4. 部署更新", 1)
    doc.add_paragraph("AI能力合并为一个agent-api服务。agent-api内部包含FastAPI、LangGraph、RAG、Tool/MCP、Model Gateway、Redis队列消费逻辑。agent-admin仅负责管理后台、权限、配置、任务下发，通过REST调用agent-api。")
    table(doc, ["服务", "端口", "职责"], [
        ["agent-admin", "8080", "Java/Spring Boot，用户、权限、配置、任务下发"],
        ["agent-api", "8000", "Python/FastAPI，Excel解析、LangGraph、RAG、Tool/MCP、Model Gateway、结果导出"],
        ["PostgreSQL + pgvector", "5432", "物料结构化库、任务、审计、向量辅助索引"],
        ["Redis", "6379", "任务队列、缓存、checkpoint辅助"],
    ])
    doc.add_heading("5. 已生成工程代码", 1)
    table(doc, ["路径", "说明"], [
        ["material_agent/excel_reader.py", "Excel读取、表头识别、workbook profile"],
        ["material_agent/pipeline.py", "PLM知识库加载、模板字段抽取、字段映射、补全输出、五步样例转换"],
        ["material_agent/cli.py", "本地验证CLI：inspect-samples/run-pilot/run-five-step"],
        ["material_agent/app.py", "FastAPI服务入口"],
        ["material_agent/requirements.txt", "工程依赖版本"],
    ])
    doc.save(ARCH_DOC)


def main():
    build_selection()
    build_arch()
    print(SELECT_DOC.resolve())
    print(ARCH_DOC.resolve())


if __name__ == "__main__":
    main()
