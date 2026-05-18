from pathlib import Path
from tempfile import TemporaryDirectory
from unittest import TestCase

import openpyxl

from material_agent.core.config import settings
from material_agent.domain.models import AgentRunRequest
from material_agent.services.completion_engine import MaterialCompletionEngine


class SampleRegressionTests(TestCase):
    def test_pilot_sample_generates_expected_workbook_shape(self):
        if not settings.plm_knowledge_file.exists() or not settings.target_template_file.exists():
            self.skipTest("Sample Excel files are not available on this machine.")

        with TemporaryDirectory() as tmp:
            output = Path(tmp) / "pilot.xlsx"
            result = MaterialCompletionEngine().run_pilot_completion(
                AgentRunRequest(
                    knowledge_file=settings.plm_knowledge_file,
                    template_file=settings.target_template_file,
                    output_file=output,
                    limit=5,
                    run_id="test-pilot",
                )
            )

            self.assertEqual(result.rows, 5)
            self.assertEqual(result.fields, 22)
            self.assertTrue(output.exists())

            workbook = openpyxl.load_workbook(output, read_only=True, data_only=True)
            self.assertEqual(workbook["补全结果"].max_row, 6)
            self.assertEqual(workbook["统计"].max_row, 5)

    def test_five_step_sample_generates_mapping_and_result(self):
        required = [
            settings.old_standard_file,
            settings.new_standard_file,
            settings.standard_mapping_file,
            settings.old_material_file,
        ]
        if not all(path.exists() for path in required):
            self.skipTest("Five-step sample Excel files are not available on this machine.")

        with TemporaryDirectory() as tmp:
            output_dir = Path(tmp)
            engine = MaterialCompletionEngine()
            mapping = engine.build_mapping_from_files(settings.old_standard_file, settings.new_standard_file)
            mapping_output = output_dir / "mapping.xlsx"
            result_output = output_dir / "result.xlsx"
            mapping.to_excel(mapping_output, index=False)

            result = engine.transform_old_material(settings.old_material_file, settings.standard_mapping_file, result_output)

            self.assertEqual(len(mapping), 19)
            self.assertEqual(result.fields, 19)
            self.assertEqual(result.rows, 1)
            self.assertTrue(mapping_output.exists())
            self.assertTrue(result_output.exists())
