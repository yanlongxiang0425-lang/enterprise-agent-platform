import os
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).parent / ".vendor"))

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from pptx import Presentation
from pptx.util import Inches as PptInches, Pt as PptPt
from pptx.dml.color import RGBColor as PptRGB
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN

from PIL import Image, ImageDraw, ImageFont


OUT = Path("交付物")
OUT.mkdir(exist_ok=True)
IMG = OUT / "images"
IMG.mkdir(exist_ok=True)

FONT_MED = "/System/Library/Fonts/STHeiti Medium.ttc"
FONT_LIGHT = "/System/Library/Fonts/STHeiti Light.ttc"


def font(size, bold=False):
    path = FONT_MED if bold else FONT_LIGHT
    return ImageFont.truetype(path, size)


def wrap_text(text, max_chars):
    lines = []
    for raw in str(text).split("\n"):
        line = ""
        for ch in raw:
            line += ch
            if len(line) >= max_chars:
                lines.append(line)
                line = ""
        if line:
            lines.append(line)
    return lines or [""]


def box(draw, xy, title, lines=None, fill="#F8FAFC", outline="#334155", title_color="#0F172A"):
    x1, y1, x2, y2 = xy
    draw.rounded_rectangle(xy, radius=18, fill=fill, outline=outline, width=3)
    draw.text((x1 + 18, y1 + 15), title, font=font(23, True), fill=title_color)
    yy = y1 + 54
    for line in lines or []:
        for w in wrap_text(line, 18):
            draw.text((x1 + 18, yy), w, font=font(18), fill="#334155")
            yy += 25


def arrow(draw, start, end, color="#64748B", width=4):
    draw.line([start, end], fill=color, width=width)
    x1, y1 = start
    x2, y2 = end
    import math
    angle = math.atan2(y2 - y1, x2 - x1)
    length = 15
    pts = []
    for delta in (math.pi * 0.82, -math.pi * 0.82):
        pts.append((x2 + length * math.cos(angle + delta), y2 + length * math.sin(angle + delta)))
    draw.polygon([(x2, y2), pts[0], pts[1]], fill=color)


def save_architecture_diagram():
    img = Image.new("RGB", (1800, 1120), "#FFFFFF")
    d = ImageDraw.Draw(img)
    d.text((60, 40), "离线 RAG Agent 总体技术架构", font=font(42, True), fill="#111827")
    d.text((60, 96), "从数据存储层到应用层：文件解析、RAG 检索、Agent 编排、LLM 路由、审计脱敏、离线部署", font=font(24), fill="#475569")

    layers = [
        ("应用层", ["Web 管理台", "Java 业务系统", "字段填充 API", "知识库管理"], "#DBEAFE", "#2563EB"),
        ("接入层", ["Nginx / APISIX", "鉴权", "限流", "审计入口"], "#E0F2FE", "#0284C7"),
        ("Agent 编排层", ["FastAPI", "LangGraph 1.x", "任务状态机", "HITL 人工确认"], "#DCFCE7", "#16A34A"),
        ("工具与 RAG 层", ["MCP/Tool Registry", "字段字典工具", "混合检索", "Rerank"], "#FEF3C7", "#D97706"),
        ("模型适配层", ["CPU: llama.cpp", "API: 百炼/方舟", "GPU: vLLM", "Embedding/Rerank"], "#F3E8FF", "#9333EA"),
        ("数据存储层", ["PostgreSQL 16", "pgvector", "excel_rows", "rag_chunks / audit"], "#FFE4E6", "#E11D48"),
    ]
    y = 170
    for name, items, fill, outline in layers:
        d.rounded_rectangle((70, y, 1730, y + 120), radius=22, fill=fill, outline=outline, width=3)
        d.text((105, y + 38), name, font=font(30, True), fill="#111827")
        x = 370
        for item in items:
            d.rounded_rectangle((x, y + 30, x + 285, y + 88), radius=16, fill="#FFFFFF", outline=outline, width=2)
            d.text((x + 18, y + 47), item, font=font(21), fill="#1F2937")
            x += 320
        if y < 170 + 5 * 145:
            arrow(d, (900, y + 120), (900, y + 145), outline, 5)
        y += 145

    d.rounded_rectangle((70, 1042, 1730, 1090), radius=16, fill="#F8FAFC", outline="#CBD5E1", width=2)
    d.text((100, 1054), "可观测横切：OpenTelemetry + Prometheus + Grafana + Phoenix；安全横切：Presidio 脱敏 + 全链路审计 + 数据域过滤", font=font(22), fill="#334155")
    path = IMG / "01_overall_architecture.png"
    img.save(path)
    return path


def save_pipeline_diagram():
    img = Image.new("RGB", (1800, 1100), "#FFFFFF")
    d = ImageDraw.Draw(img)
    d.text((60, 40), "离线构建与在线查询双 Pipeline", font=font(42, True), fill="#111827")
    d.text((60, 96), "重点解决 Excel 行级 chunk、字段填充、指定字段精确查询与 pgvector 语义召回的组合问题", font=font(24), fill="#475569")

    d.rounded_rectangle((50, 150, 1750, 555), radius=28, fill="#F8FAFC", outline="#22C55E", width=4)
    d.text((80, 178), "A. 离线构建 / 写入 RAG", font=font(30, True), fill="#166534")
    stages = [
        ("文件输入", ["Excel/CSV", "Word/PPT/PDF", "TXT/Markdown"]),
        ("解析标准化", ["Excel 结构化解析", "文档段落/页码/OCR", "字段归一化"]),
        ("分片策略", ["Excel: 一行=主chunk", "row_json保存整行", "长文本二级分片"]),
        ("向量与索引", ["bge-m3 embedding", "B-tree/GIN/trgm", "pgvector HNSW"]),
        ("质量与回滚", ["批次状态", "抽样召回", "失败批次回滚"]),
    ]
    x = 90
    for i, (t, lines) in enumerate(stages):
        box(d, (x, 250, x + 270, 465), t, lines, "#ECFDF5", "#16A34A")
        if i < len(stages) - 1:
            arrow(d, (x + 270, 357), (x + 325, 357), "#16A34A")
        x += 335

    d.rounded_rectangle((50, 625, 1750, 1030), radius=28, fill="#F8FAFC", outline="#7C3AED", width=4)
    d.text((80, 653), "B. 在线查询 / 字段填充 RAG", font=font(30, True), fill="#5B21B6")
    stages2 = [
        ("用户问题", ["字段填充", "事实问答", "模糊查询"]),
        ("查询理解", ["意图识别", "字段名/值抽取", "字段别名映射"]),
        ("检索路由", ["明确字段→精确查", "自然语言→向量查", "混合→过滤后召回"]),
        ("结果融合", ["row_id优先", "TopK重排", "置信度评分"]),
        ("答案输出", ["字段值直返", "LLM解释", "引用来源/审计"]),
    ]
    x = 90
    for i, (t, lines) in enumerate(stages2):
        box(d, (x, 725, x + 270, 940), t, lines, "#F5F3FF", "#7C3AED")
        if i < len(stages2) - 1:
            arrow(d, (x + 270, 832), (x + 325, 832), "#7C3AED")
        x += 335

    path = IMG / "02_pipeline.png"
    img.save(path)
    return path


def save_sequence_diagram():
    img = Image.new("RGB", (1800, 1000), "#FFFFFF")
    d = ImageDraw.Draw(img)
    d.text((60, 40), "字段填充查询时序图", font=font(42, True), fill="#111827")
    actors = ["用户/Java系统", "API网关", "FastAPI", "LangGraph Agent", "RAG服务", "PostgreSQL/pgvector", "LLM Router"]
    x_positions = [120, 360, 600, 860, 1120, 1380, 1620]
    top = 150
    for x, a in zip(x_positions, actors):
        d.rounded_rectangle((x - 95, top, x + 95, top + 60), radius=14, fill="#E0F2FE", outline="#0284C7", width=2)
        d.text((x - 72, top + 18), a, font=font(18, True), fill="#0F172A")
        d.line((x, top + 70, x, 900), fill="#CBD5E1", width=3)
    steps = [
        (0, 1, "1. 提交问题/字段填充请求"),
        (1, 2, "2. 鉴权、限流、数据域过滤"),
        (2, 3, "3. 创建任务状态 + 脱敏"),
        (3, 4, "4. Tool: 查询理解与检索路由"),
        (4, 5, "5a. 明确字段值：精确查询 excel_rows"),
        (4, 5, "5b. 自然语言：pgvector 语义召回"),
        (4, 3, "6. 返回 row_id/chunk/source"),
        (3, 6, "7. 复杂问答调用本地/API LLM"),
        (6, 3, "8. JSON Schema 约束输出"),
        (3, 2, "9. 结果融合、置信度、审计"),
        (2, 0, "10. 返回字段值/答案/引用"),
    ]
    y = 250
    for frm, to, label in steps:
        x1, x2 = x_positions[frm], x_positions[to]
        color = "#16A34A" if "5a" in label else "#7C3AED" if "5b" in label or "LLM" in label else "#475569"
        arrow(d, (x1, y), (x2, y), color, 4)
        d.text((min(x1, x2) + 12, y - 28), label, font=font(17), fill="#111827")
        y += 58
    path = IMG / "03_sequence.png"
    img.save(path)
    return path


def save_deployment_diagram():
    img = Image.new("RGB", (1800, 980), "#FFFFFF")
    d = ImageDraw.Draw(img)
    d.text((60, 40), "单机虚拟机离线部署架构", font=font(42, True), fill="#111827")
    d.text((60, 96), "前期 CPU 验证功能，模型确定后再采购 GPU；组件以 Docker Compose/离线包方式交付。", font=font(24), fill="#475569")
    d.rounded_rectangle((70, 160, 1730, 870), radius=28, fill="#F8FAFC", outline="#334155", width=4)
    d.text((100, 190), "内网虚拟机 / 单机节点", font=font(30, True), fill="#111827")
    boxes = [
        ((120, 270, 430, 410), "nginx / APISIX", ["80/443", "反向代理、限流、鉴权"], "#DBEAFE", "#2563EB"),
        ((520, 270, 850, 410), "rag-agent-api", ["FastAPI + LangGraph 1.x", "Agent编排、工具调用"], "#DCFCE7", "#16A34A"),
        ((940, 270, 1270, 410), "worker", ["Celery/RQ", "文档解析、embedding、批处理"], "#FEF3C7", "#D97706"),
        ((1360, 270, 1660, 410), "llm-router", ["CPU: llama.cpp", "API/GPU模式可切换"], "#F3E8FF", "#7C3AED"),
        ((120, 520, 430, 680), "PostgreSQL 16", ["pgvector、JSONB", "rag_chunks / excel_rows / audit"], "#FFE4E6", "#E11D48"),
        ((520, 520, 850, 680), "Redis 7", ["缓存、队列、限流计数", "schema hash / query cache"], "#FFEDD5", "#EA580C"),
        ((940, 520, 1270, 680), "文件与模型目录", ["原始文件、离线模型", "wheelhouse / 镜像 tar"], "#E0F2FE", "#0284C7"),
        ((1360, 520, 1660, 680), "可观测", ["Prometheus/Grafana", "Phoenix/Loki 可选"], "#ECFDF5", "#059669"),
    ]
    for xy, title, lines, fill, outline in boxes:
        box(d, xy, title, lines, fill, outline)
    for s, e in [((430, 340), (520, 340)), ((850, 340), (940, 340)), ((1270, 340), (1360, 340)), ((685, 410), (685, 520)), ((1105, 410), (1105, 520)), ((685, 680), (275, 680)), ((1105, 680), (1510, 680))]:
        arrow(d, s, e, "#475569", 4)
    d.text((120, 770), "最低 CPU 验证配置：16C / 64GB / 500GB SSD，可跑通解析、RAG、Agent、审计；7B GGUF 模型只做功能验证，延迟不作为性能指标。", font=font(23), fill="#334155")
    d.text((120, 812), "GPU 生产建议：32C / 128GB / 2TB SSD + L20 48GB 或 A100 40GB；vLLM 承载 14B/32B 量化模型。", font=font(23), fill="#334155")
    path = IMG / "04_deployment.png"
    img.save(path)
    return path


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def doc_styles(doc):
    styles = doc.styles
    styles["Normal"].font.name = "Arial Unicode MS"
    styles["Normal"].font.size = Pt(10.5)
    for s in ["Heading 1", "Heading 2", "Heading 3"]:
        styles[s].font.name = "Arial Unicode MS"
        styles[s].font.color.rgb = RGBColor(17, 24, 39)
    styles["Heading 1"].font.size = Pt(18)
    styles["Heading 2"].font.size = Pt(15)
    styles["Heading 3"].font.size = Pt(12)


def add_title(doc, title, subtitle):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(title)
    r.bold = True
    r.font.size = Pt(22)
    r.font.name = "Arial Unicode MS"
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(subtitle)
    r.font.size = Pt(11)
    r.font.color.rgb = RGBColor(71, 85, 105)
    doc.add_paragraph("")


def add_bullets(doc, items):
    for item in items:
        doc.add_paragraph(item, style="List Bullet")


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = h
        set_cell_shading(hdr[i], "E2E8F0")
        hdr[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for p in hdr[i].paragraphs:
            for r in p.runs:
                r.bold = True
    for row in rows:
        cells = table.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = str(val)
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
    doc.add_paragraph("")
    return table


def add_image(doc, path, caption):
    doc.add_picture(str(path), width=Inches(6.8))
    last = doc.paragraphs[-1]
    last.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = doc.add_paragraph(caption)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for r in p.runs:
        r.font.size = Pt(9)
        r.font.color.rgb = RGBColor(100, 116, 139)


def build_tech_doc(paths):
    doc = Document()
    doc_styles(doc)
    add_title(doc, "离线 RAG Agent 技术选型文档", "面向 Excel/Word/PPT/PDF 知识库、字段填充与企业内网离线部署场景")

    doc.add_heading("1. 需求复述", 1)
    add_bullets(doc, [
        "建设一个从 Dify 迁出的全新 RAG Agent 架构，覆盖离线知识构建、在线查询、字段填充、审计脱敏与内网部署。",
        "数据来源以 Excel 表格为主，同时兼容 Word、PPT、PDF、TXT/Markdown/HTML。Excel 需重点支持“给定某字段值，填充同一行其他字段”的结构化查询场景。",
        "知识库规模不大，字段量级几万，优先沿用 PostgreSQL + pgvector；当前阶段不优先引入 Elasticsearch。",
        "部署环境为内网单机虚拟机，不能连接外网，需要离线 wheel 包、Docker 镜像包、模型文件包。",
        "模型策略先以 CPU 跑通功能为准，确认模型效果后再选择 GPU；LangChain/LangGraph 必须采用 1.0 以上版本。",
    ])

    doc.add_heading("2. 总体选型结论", 1)
    add_table(doc, ["类别", "推荐组件与版本", "选择原因"], [
        ("主语言", "Python 3.11.9", "AI/RAG/Agent 生态成熟，LangChain、LangGraph、vLLM、Embedding、Rerank、文档解析均以 Python 支持最完整。"),
        ("业务接入", "Java 17 / Spring Boot 3.3.x", "公司主栈可保留在网关、业务系统、权限与管理后台层，通过 HTTP/gRPC 调用 Python AI 服务。"),
        ("Agent 编排", "LangGraph >= 1.0.3，LangChain >= 1.0.5", "状态机编排、节点可测试、可回放、可中断恢复，适合替代 Dify Workflow。"),
        ("API 服务", "FastAPI 0.115.x + Uvicorn 0.32.x", "轻量高性能，Pydantic 2.x 强类型校验，适合 SSE 流式输出。"),
        ("任务队列", "Celery 5.4.x / RQ 2.x + Redis 7.4.x", "离线文档解析、embedding、批量导入、失败重试。"),
        ("向量库", "PostgreSQL 16.x + pgvector 0.8.x", "当前规模下足够；同时支持普通列索引、JSONB GIN、pg_trgm 与向量索引。"),
        ("Embedding", "BAAI bge-m3", "中文/多语言表现稳定，支持 dense/sparse/ColBERT 多形态，适合混合检索。"),
        ("Rerank", "bge-reranker-v2-m3", "对 top-k 召回结果做二次排序，提升引用准确率。"),
        ("CPU 验证模型", "Qwen2.5-7B-Instruct GGUF Q4_K_M + llama.cpp", "先验证链路与 Prompt，不把 CPU 延迟作为最终性能指标。"),
        ("GPU 生产模型", "Qwen2.5-14B/32B-Instruct-AWQ 或 Qwen3-30B-A3B", "在成本、效果、吞吐之间平衡；效果确认后再采购 GPU。"),
        ("推理服务", "CPU: llama.cpp / GPU: vLLM 0.8.x+", "同一 LLM Router 暴露 OpenAI 兼容接口，后续切换模型不改业务代码。"),
        ("脱敏审计", "Microsoft Presidio 2.2.x + 自定义正则 + PostgreSQL audit schema", "手机号、身份证、邮箱、合同号等敏感字段可脱敏并保留审计追踪。"),
        ("可观测", "OpenTelemetry 1.30.x + Prometheus + Grafana + Phoenix", "指标、日志、Trace、LLM 调用链路统一观测，Phoenix 离线友好。"),
    ])

    doc.add_heading("3. Java 与 Python 技术栈对比", 1)
    add_table(doc, ["维度", "Java 技术栈", "Python 技术栈", "建议"], [
        ("公司基础", "团队熟悉，工程治理、权限、网关、后台优势明显", "AI 团队需补工程规范", "Java 保留在业务接入层"),
        ("AI 生态", "LangChain4j、Spring AI 可用，但很多新能力滞后", "LangChain/LangGraph/vLLM/Transformers 首发支持", "Agent/RAG 核心选 Python"),
        ("文档解析", "Apache POI、Tika 稳定", "openpyxl、pandas、unstructured、pymupdf 组合灵活", "Excel 可用 Python，Java 可负责上传与权限"),
        ("模型推理", "本地推理生态弱，通常需要调用服务", "vLLM、llama.cpp、sentence-transformers 完整", "模型服务选 Python"),
        ("运维治理", "Spring Boot 体系成熟", "需明确日志、配置、类型与测试规范", "用 Docker Compose + 规范化配置补齐"),
        ("综合结论", "适合做企业集成与管理面", "适合做 AI 生产力核心", "Java + Python 双服务边界最合理"),
    ])

    doc.add_heading("4. pgvector 是否满足", 1)
    add_bullets(doc, [
        "当前文档数、字段数和 chunk 数量级较小，PostgreSQL + pgvector 完全可以覆盖第一阶段需求。",
        "Excel 场景的核心不是全文搜索，而是“字段精确定位 + 行级 row_json 回填 + 必要时语义补充”。这正是 PostgreSQL 的优势。",
        "建议建立 rag_chunks、excel_rows、field_dictionary、ingest_batches、audit_logs 五类核心表。",
        "索引组合：常用字段 B-tree；row_json 用 JSONB GIN；模糊字段用 pg_trgm；语义召回用 pgvector HNSW。",
        "只有在非结构化全文检索规模明显扩大、中文复杂分词/高亮/聚合成为刚需、或 PostgreSQL 性能压测不达标时，再引入 Elasticsearch/OpenSearch。",
    ])
    add_image(doc, paths["pipeline"], "图 1：离线构建与在线查询双 Pipeline")

    doc.add_heading("5. 最低环境配置", 1)
    add_table(doc, ["阶段", "最低配置", "用途", "说明"], [
        ("CPU 功能验证", "16C / 64GB RAM / 500GB SSD", "跑通上传、解析、入库、检索、Agent、审计", "7B GGUF 可慢速验证；不做性能承诺。"),
        ("CPU 增强验证", "32C / 128GB RAM / 1TB SSD", "多文档批量导入、Rerank、并发联调", "适合验证 100 QPS API 层与非模型环节。"),
        ("GPU 生产起步", "32C / 128GB RAM / 2TB SSD + L20 48GB 或 A100 40GB", "14B/32B 量化模型 vLLM 推理", "确认模型后再采购，避免先买错卡。"),
        ("离线包", "Docker 镜像 tar + pip wheelhouse + 模型权重 + 初始化 SQL", "内网无外网部署", "建议统一版本锁定和 SHA256 校验。"),
    ])

    doc.add_heading("6. 离线包清单", 1)
    add_table(doc, ["包类型", "内容", "估算体积"], [
        ("Python wheelhouse", "fastapi、langchain、langgraph、psycopg、sqlalchemy、pandas、openpyxl、pymupdf、presidio、opentelemetry 等", "2-5GB"),
        ("Docker 镜像", "postgres、redis、nginx/apisix、rag-agent-api、worker、phoenix/prometheus/grafana 可选", "8-20GB"),
        ("模型文件", "bge-m3、bge-reranker-v2-m3、Qwen2.5-7B GGUF、候选 14B/32B 权重", "20-100GB"),
        ("数据库脚本", "schema、索引、初始化字典、审计表、示例数据", "<100MB"),
    ])

    doc.add_heading("7. 参考来源", 1)
    add_bullets(doc, [
        "LangChain 官方文档：https://docs.langchain.com/",
        "LangGraph 官方文档：https://docs.langchain.com/oss/python/langgraph/overview",
        "pgvector 项目：https://github.com/pgvector/pgvector",
    ])
    doc.save(OUT / "离线RAG智能体_技术选型文档.docx")


def build_arch_doc(paths):
    doc = Document()
    doc_styles(doc)
    add_title(doc, "离线 RAG Agent 架构设计文档", "覆盖离线写入、在线查询、字段填充、Agent/Tool/MCP 交互、部署与时序")

    doc.add_heading("1. 需求与场景理解", 1)
    add_bullets(doc, [
        "本系统不是单纯聊天机器人，而是面向企业内网文档与 Excel 表格的离线 RAG Agent。",
        "第一类核心场景：离线构建知识库，把 Excel/Word/PPT/PDF/TXT 等文件解析、分片、向量化并写入 PostgreSQL + pgvector。",
        "第二类核心场景：用户查询或字段填充。若用户给出明确字段值，系统优先精准定位 Excel 行，再返回同一行其他字段；若是自然语言问题，再走语义召回和 LLM 生成。",
        "系统需支持审计、脱敏、人工确认、错误样本回流、内网离线部署。"
    ])

    doc.add_heading("2. 总体技术架构", 1)
    add_image(doc, paths["arch"], "图 1：离线 RAG Agent 总体技术架构")
    doc.add_paragraph("系统按应用层、接入层、Agent 编排层、工具与 RAG 层、模型适配层、数据存储层分层。Java 业务系统通过网关调用 Python AI 服务；Python 内部使用 LangGraph 1.x 编排状态机，将文档解析、查询理解、检索路由、工具调用、LLM 输出约束、人工确认串成可追踪流程。")

    doc.add_heading("3. 离线构建与在线查询流程", 1)
    add_image(doc, paths["pipeline"], "图 2：离线构建与在线查询双 Pipeline")
    doc.add_heading("3.1 离线写入流程", 2)
    add_bullets(doc, [
        "文件导入：接收 Excel/CSV、Word、PPT、PDF、TXT/Markdown/HTML，并生成 source_id、batch_id、checksum。",
        "类型路由：Excel 使用结构化解析；Word/PPT/PDF 按标题、页码、slide、OCR 结果解析。",
        "Excel 行级 chunk：一行业务记录作为主 chunk，chunk_text 用“字段名: 字段值”串联，row_json 保存整行字段。",
        "索引写入：rag_chunks 写向量与 chunk_text，excel_rows 写 row_json 与常用字段列，field_dictionary 写字段别名和语义解释。",
        "质量检查：检查空字段、重复行、chunk 数、embedding 失败、抽样召回效果；失败批次可回滚。"
    ])
    doc.add_heading("3.2 在线查询流程", 2)
    add_bullets(doc, [
        "查询理解：识别字段填充、事实问答、模糊查询三类意图；抽取字段名、字段值与目标字段。",
        "字段别名映射：例如“合同号 = 合同编号”“客户 = 客户名称”。字段字典可由人工维护，也可从 Excel 表头自动生成候选。",
        "检索路由：有明确字段值时优先 PostgreSQL 精确查询；纯自然语言走 pgvector；二者都有时先字段过滤再向量召回。",
        "结果融合：精确 row_id 命中优先，语义 chunk 补充上下文，rerank 重新排序，输出置信度。",
        "字段填充：若只需返回同一行其他字段，可不调用 LLM，直接返回结构化结果，降低成本和幻觉风险。"
    ])

    doc.add_heading("4. 字段填充查询时序", 1)
    add_image(doc, paths["sequence"], "图 3：字段填充查询时序图")

    doc.add_heading("5. Agent、Tool、MCP、Skill 交互细节", 1)
    add_table(doc, ["模块", "职责", "交互方式"], [
        ("LangGraph Agent", "维护任务状态：ingest/query/route/retrieve/answer/audit/hitl", "内部节点函数 + 状态对象；支持断点恢复。"),
        ("Tool Registry", "注册字段字典、RAG 检索、精确查询、审计写入、脱敏等工具", "Agent 通过工具名与结构化参数调用。"),
        ("MCP Server（可选）", "把数据库查询、文件系统、内部知识源封装成标准工具", "适合后期跨系统复用工具能力。"),
        ("Skill（可选）", "沉淀 Excel 行级分片、字段别名学习、质量检查等专家流程", "作为 Agent 调用前后的确定性流程，不让 LLM 处理确定性逻辑。"),
        ("LLM Router", "在 CPU/API/GPU 三模式之间路由，统一 OpenAI 兼容接口", "通过配置切换，不影响上层 Agent。"),
    ])

    doc.add_heading("6. 数据模型设计", 1)
    add_table(doc, ["表", "关键字段", "说明"], [
        ("rag_chunks", "chunk_id, source_id, chunk_text, embedding, metadata, row_id", "RAG 召回主表，metadata 保存页码、sheet、字段等。"),
        ("excel_rows", "row_id, source_id, sheet_name, row_no, row_json, common_fields", "字段填充主表，保存 Excel 原始行级数据。"),
        ("field_dictionary", "field_name, aliases, data_type, source_scope, examples", "字段别名、语义解释、枚举值，用于查询理解。"),
        ("ingest_batches", "batch_id, status, counts, error_message, checksum", "离线导入批次状态、失败回滚与重建。"),
        ("audit_logs", "request_id, user_id, query, filters, hit_row_ids, answer, masked_payload", "审计与问题追踪。"),
    ])

    doc.add_heading("7. 部署架构", 1)
    add_image(doc, paths["deployment"], "图 4：单机虚拟机离线部署架构")

    doc.add_heading("8. 性能与高可用策略", 1)
    add_bullets(doc, [
        "大 JSON 不直接喂给模型：先解析入库，再按字段/业务域/表集合分批处理，避免上下文爆炸。",
        "结构化优先：字段填充走 PostgreSQL 精确查询，不必要时不调用 LLM。",
        "缓存：schema hash、字段别名、query embedding、常见查询结果放 Redis。",
        "批处理：embedding 与文档解析走 worker 队列，失败可重试，批次可回滚。",
        "限流：网关限流 + Redis 计数；LLM Router 按模型能力做并发控制。",
        "可观测：每个 request_id 串联 API、Agent、Tool、DB、LLM 调用链路。"
    ])
    doc.save(OUT / "离线RAG智能体_架构设计文档.docx")


def add_ppt_title(slide, text, subtitle=None):
    box = slide.shapes.add_textbox(PptInches(0.55), PptInches(0.35), PptInches(12.2), PptInches(0.55))
    p = box.text_frame.paragraphs[0]
    p.text = text
    p.font.size = PptPt(26)
    p.font.bold = True
    p.font.color.rgb = PptRGB(17, 24, 39)
    if subtitle:
        s = slide.shapes.add_textbox(PptInches(0.58), PptInches(0.92), PptInches(12), PptInches(0.35))
        sp = s.text_frame.paragraphs[0]
        sp.text = subtitle
        sp.font.size = PptPt(12)
        sp.font.color.rgb = PptRGB(71, 85, 105)


def add_ppt_bullets(slide, items, x=0.75, y=1.45, w=5.7, h=4.8):
    tb = slide.shapes.add_textbox(PptInches(x), PptInches(y), PptInches(w), PptInches(h))
    tf = tb.text_frame
    tf.clear()
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = item
        p.font.size = PptPt(15)
        p.font.color.rgb = PptRGB(30, 41, 59)
        p.space_after = PptPt(6)
        p.level = 0


def add_ppt_image(slide, path, x=0.55, y=1.25, w=12.2):
    slide.shapes.add_picture(str(path), PptInches(x), PptInches(y), width=PptInches(w))


def build_ppt(paths):
    prs = Presentation()
    prs.slide_width = PptInches(13.333)
    prs.slide_height = PptInches(7.5)
    blank = prs.slide_layouts[6]

    slide = prs.slides.add_slide(blank)
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = PptRGB(15, 23, 42)
    t = slide.shapes.add_textbox(PptInches(0.8), PptInches(1.5), PptInches(11.6), PptInches(1.2))
    p = t.text_frame.paragraphs[0]
    p.text = "离线 RAG Agent 架构与技术选型"
    p.font.size = PptPt(38)
    p.font.bold = True
    p.font.color.rgb = PptRGB(255, 255, 255)
    s = slide.shapes.add_textbox(PptInches(0.85), PptInches(2.7), PptInches(10.8), PptInches(0.6))
    sp = s.text_frame.paragraphs[0]
    sp.text = "Excel 行级字段填充 · PostgreSQL + pgvector · LangGraph 1.x · 内网离线部署"
    sp.font.size = PptPt(18)
    sp.font.color.rgb = PptRGB(203, 213, 225)

    slide = prs.slides.add_slide(blank)
    add_ppt_title(slide, "需求复述", "从 Dify/旧方案抽象出的两条核心链路")
    add_ppt_bullets(slide, [
        "离线构建：多类型文件解析、Excel 行级 chunk、embedding、写入 PostgreSQL + pgvector。",
        "在线查询：字段值精确过滤、语义召回、结果融合、字段填充或 LLM 生成。",
        "内网单机虚机部署，无外网；先 CPU 验证功能，确认模型后再选 GPU。",
        "需要审计、脱敏、可观测、人工确认与错误样本回流。",
    ], w=12)

    for title, img in [
        ("总体技术架构", paths["arch"]),
        ("离线构建与在线查询流程", paths["pipeline"]),
        ("字段填充查询时序", paths["sequence"]),
        ("单机虚拟机部署架构", paths["deployment"]),
    ]:
        slide = prs.slides.add_slide(blank)
        add_ppt_title(slide, title)
        add_ppt_image(slide, img, y=1.1, w=12.1)

    slide = prs.slides.add_slide(blank)
    add_ppt_title(slide, "技术选型结论")
    add_ppt_bullets(slide, [
        "AI 核心：Python 3.11 + FastAPI + LangGraph 1.x + LangChain 1.x。",
        "业务接入：Java 17 / Spring Boot 3.3.x 继续承担网关、权限、后台与企业集成。",
        "存储：PostgreSQL 16 + pgvector 0.8；结构化字段优先，向量召回补充。",
        "模型：CPU 阶段 llama.cpp + 7B GGUF；GPU 阶段 vLLM + 14B/32B 量化模型。",
        "离线交付：Docker 镜像 tar、pip wheelhouse、模型权重、初始化 SQL、验收脚本。",
    ], w=12)

    slide = prs.slides.add_slide(blank)
    add_ppt_title(slide, "为什么第一阶段不引入 Elasticsearch")
    add_ppt_bullets(slide, [
        "当前主需求是 Excel 指定字段查找、同一行字段补全，不是大规模全文搜索。",
        "PostgreSQL 同时具备 B-tree、JSONB GIN、pg_trgm、全文检索和 pgvector。",
        "减少组件数量更适合内网离线部署，也降低运维和故障定位成本。",
        "当非结构化全文规模、中文复杂分词、高亮、聚合统计、高并发搜索成为刚需，再引入 ES/OpenSearch。",
    ], w=12)

    prs.save(OUT / "离线RAG智能体_架构与选型汇报.pptx")


def main():
    paths = {
        "arch": save_architecture_diagram(),
        "pipeline": save_pipeline_diagram(),
        "sequence": save_sequence_diagram(),
        "deployment": save_deployment_diagram(),
    }
    build_tech_doc(paths)
    build_arch_doc(paths)
    build_ppt(paths)
    print("generated:")
    for p in sorted(OUT.iterdir()):
        if p.suffix.lower() in {".docx", ".pptx"}:
            print(p.resolve())


if __name__ == "__main__":
    main()
